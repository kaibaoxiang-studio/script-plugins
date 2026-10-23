#!/usr/bin/env python3
"""道具异称、交接、同名不同物及场次归属回归测试。"""
import copy
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook, load_workbook
from prop_identity import prop_candidates, related_props, sync_props, validate_props
from build_xlsx import validate, write_asset_sheet, write_breakdown


def fixture():
    texts = ['凯尔拔出屠龙匕首，刀柄刻有三道血焰纹。', '凯尔将这把匕首交给莱拉。',
             '莱拉举起龙骨匕首，刀柄仍是三道血焰纹。', '她将它折断。',
             '旁白只提到屠龙匕首，画面没有展示。']
    scenes = [{'scene_id': f'1-{i}', 'location': '大厅', 'int_ext': '内', 'time': '日',
               'characters': [], 'lines': [{'type': 'action' if i < 5 else 'dialog', 'text': text,
                                           **({'speaker': '旁白', 'voice': 'VO'} if i == 5 else {})}]}
              for i, text in enumerate(texts, 1)]
    data = {'episodes': [{'episode': 1, 'title': '测试', 'scenes': scenes}]}

    def mention(i, raw, basis, state='S01', anchor=None):
        text, low = texts[i - 1], texts[i - 1].index(raw)
        result = {'id': f'M00{i}', 'source': {'scene_id': f'1-{i}', 'line': 1, 'quote': text},
                  'raw_name': raw, 'start_offset': low, 'end_offset': low + len(raw),
                  'basis': basis, 'timeline': '连续段1', 'visibility': 'visible' if i < 5 else 'reference',
                  'state_id': state if i < 5 else None, 'holder': '凯尔' if i == 1 else '莱拉',
                  'note': '' if i < 5 else '仅画外音提及，不计画面场次。'}
        if anchor:
            result.update({'from': anchor['id'], 'identity_refs': [copy.deepcopy(anchor['source'])],
                           'identity_reason': '沿用实物交接链，本连续段未出现另一把匕首。'})
        return result
    m1 = mention(1, '屠龙匕首', 'named')
    m2 = mention(2, '这把匕首', 'contextual', anchor=m1)
    m3 = mention(3, '龙骨匕首', 'named')
    m4 = mention(4, '它', 'contextual', 'S02', m3)
    m5 = mention(5, '屠龙匕首', 'named')
    prop = {'id': 'PRP-001', 'name': '屠龙匕首', 'aliases': ['龙骨匕首'], 'owner': '',
            'description': '刀柄刻三道血焰纹。', 'note': '',
            'alias_evidence': [{'alias': '龙骨匕首', 'basis': 'distinctive',
                                'reason': '交接前后均有三道血焰纹，指代链连续。',
                                'evidence_refs': [copy.deepcopy(m1['source']), copy.deepcopy(m3['source'])]}],
            'mentions': [m1, m2, m3, m4, m5],
            'states': [{'id': 'S01', 'name': '完整', 'scenes': []}, {'id': 'S02', 'name': '折断', 'scenes': []}]}
    assets = {'characters': [], 'scenes': [{'id': 'SCN-001', 'name': '大厅', 'int_ext': '内', 'times': ['日'],
                                          'states': [{'id': 'S01', 'name': '常态', 'scenes': [s['scene_id'] for s in scenes]}]}],
              'props': [prop]}
    return data, assets


class PropIdentityTests(unittest.TestCase):
    def setUp(self):
        self.data, self.assets = fixture()
        self.prop = self.assets['props'][0]

    def sync(self):
        self.assertEqual(sync_props(self.data, self.assets), ([], []))

    def rejected(self, phrase):
        errors, _ = validate_props(self.data, self.assets, check_derived=False)
        self.assertTrue(any(phrase in error for error in errors), errors)

    def test_alias_transfer_breakage_and_reference_share_one_identity(self):
        self.sync()
        self.assertEqual(validate(self.data, self.assets), ([], []))
        self.assertEqual(self.prop['states'][0]['scenes'], ['1-1', '1-2', '1-3'])
        self.assertEqual(self.prop['states'][1]['scenes'], ['1-4'])
        self.assertEqual(len(self.assets['props']), 1)

    def test_alias_without_evidence_rejected(self):
        self.prop['alias_evidence'] = []
        self.rejected('缺少 alias_evidence')

    def test_generic_name_cannot_be_global_alias(self):
        self.prop['aliases'].append('匕首')
        self.prop['alias_evidence'].append({'alias': '匕首', 'basis': 'explicit', 'reason': '测试',
                                           'evidence_refs': [self.prop['mentions'][0]['source']]})
        self.rejected('泛称「匕首」不能作全局')

    def test_contextual_reference_requires_same_identity_anchor(self):
        self.prop['mentions'][1]['from'] = 'M999'
        self.rejected('须指向同实物同连续段')

    def test_contextual_reference_cannot_cross_timeline(self):
        self.prop['mentions'][1]['timeline'] = '多年后'
        self.rejected('须指向同实物同连续段')

    def test_evidence_must_cover_canonical_identity_and_alias(self):
        self.prop['alias_evidence'][0]['evidence_refs'] = [self.prop['mentions'][2]['source'], self.prop['mentions'][3]['source']]
        self.rejected('异称判读须覆盖')

    def test_false_alias_quote_rejected(self):
        self.prop['alias_evidence'][0]['evidence_refs'][0]['quote'] = '龙骨匕首就是屠龙匕首。'
        self.rejected('引文不在指定原文中')

    def test_duplicate_asset_alias_is_not_silently_accepted(self):
        other = copy.deepcopy(self.prop)
        other.update(id='PRP-002', name='龙骨匕首', aliases=[], alias_evidence=[])
        self.assets['props'].append(other)
        self.rejected('共用全局名称')

    def test_wrong_scene_list_cannot_pass(self):
        self.sync()
        self.prop['states'][0]['scenes'].append('1-5')
        errors, _ = validate(self.data, self.assets)
        self.assertTrue(any('scenes 与实物指代记录不一致' in error for error in errors), errors)

    def test_dialogue_does_not_count_as_visible_prop(self):
        self.prop['mentions'][-1].update(visibility='visible', state_id='S01')
        self.rejected('台词提及不能直接当作道具出镜')

    def test_mentions_have_to_match_actual_text_positions(self):
        self.prop['mentions'][0]['start_offset'] += 1
        self.rejected('原文称呼与字符位置不一致')

    def test_legacy_data_needs_identity_review(self):
        del self.prop['mentions']
        self.rejected('缺少 mentions')

    def test_longer_name_wins_over_contained_generic_word(self):
        first = [m for m in prop_candidates(self.data, self.prop, ['匕首']) if m['source']['scene_id'] == '1-1']
        self.assertEqual([m['raw_name'] for m in first], ['屠龙匕首'])

    def test_similar_type_is_candidate_not_auto_merge(self):
        self.assets['props'].append({'id': 'PRP-002', 'name': '银匕首', 'aliases': []})
        before = copy.deepcopy(self.assets)
        self.assertEqual(list(related_props(self.assets))[0]['decision'], '待原文判读')
        self.assertEqual(self.assets, before)

    def test_bad_identity_field_types_return_errors(self):
        self.prop['alias_evidence'][0]['evidence_refs'] = 7
        errors, _ = validate(self.data, self.assets)
        self.assertTrue(any('类型错误' in error for error in errors), errors)

    def test_confirmed_name_occurrence_cannot_be_omitted(self):
        self.prop['mentions'].pop()
        self.rejected('原文名称未核对: 1-5')

    def test_qualified_display_name_keeps_original_name_coverage(self):
        self.prop['name'] = '屠龙匕首·凯尔'
        self.prop['mentions'][0].update(basis='introduced', identity_reason='原文引入凯尔的屠龙匕首。')
        self.prop['mentions'][-1].update(basis='introduced', identity_reason='原文仅提及该实物。')
        self.sync()
        self.prop['mentions'].pop()
        self.rejected('原文名称未核对: 1-5')

    def test_source_name_separate_from_display_name(self):
        self.prop.update(name='屠龙匕首·凯尔', source_name='屠龙匕首')
        self.sync()
        self.assertEqual(validate(self.data, self.assets), ([], []))
        wb = Workbook()
        write_asset_sheet(wb, 'props', '道具', '道具名', [('owner', '归属人物')], self.assets, self.data)
        self.assertEqual(wb['道具'].cell(2, 2).value, '屠龙匕首·凯尔')
        self.prop['mentions'].pop()
        self.rejected('原文名称未核对: 1-5')

    def test_source_name_requires_actual_original_record(self):
        self.prop['source_name'] = '不存在的称呼'
        self.rejected('source_name 缺少实际原文')

    def test_source_name_type_is_validated(self):
        self.prop['source_name'] = ['屠龙匕首']
        self.rejected('source_name 类型错误')

    def test_candidates_search_source_name_without_existing_mentions(self):
        self.prop.update(name='屠龙匕首·凯尔', source_name='屠龙匕首', mentions=[])
        found = list(prop_candidates(self.data, self.prop))
        self.assertTrue(any(m['source']['scene_id'] == '1-5' and m['raw_name'] == '屠龙匕首' for m in found))

    def test_same_line_alias_transfer_uses_distinct_character_positions(self):
        text = '凯尔将屠龙匕首交给莱拉，莱拉接过龙骨匕首，刀柄刻有三道血焰纹。'
        scene = self.data['episodes'][0]['scenes'][0]
        scene['lines'][0]['text'] = text
        m1, m3 = self.prop['mentions'][0], self.prop['mentions'][2]
        for m in [m1, m3]:
            m['source'] = {'scene_id': '1-1', 'line': 1, 'quote': text}
            m.update(start_offset=text.index(m['raw_name']), end_offset=text.index(m['raw_name']) + len(m['raw_name']))
        self.prop['mentions'] = [m1, m3]
        self.prop['states'] = self.prop['states'][:1]
        self.data['episodes'][0]['scenes'] = [scene]
        self.assets['scenes'][0]['states'][0]['scenes'] = ['1-1']
        for basis in ['continuity', 'distinctive']:
            self.prop['alias_evidence'][0].update(basis=basis, evidence_refs=[copy.deepcopy(m1['source'])])
            self.sync()
            self.assertEqual(validate(self.data, self.assets), ([], []))

    def test_alias_evidence_quote_must_cover_actual_names(self):
        self.prop['alias_evidence'][0]['evidence_refs'][0]['quote'] = '三道血焰纹'
        self.rejected('异称判读须覆盖')

    def test_missing_or_empty_states_returns_error_without_mutation(self):
        for value in [None, []]:
            data, assets = fixture()
            prop = assets['props'][0]
            for m in prop['mentions']:
                m.update(visibility='reference', state_id=None, note='仅提及，不展示。')
            if value is None:
                del prop['states']
            else:
                prop['states'] = value
            before = copy.deepcopy(assets)
            errors, _ = sync_props(data, assets)
            self.assertTrue(any('缺少非空 states' in e for e in errors), errors)
            self.assertEqual(assets, before)

    def test_cli_missing_states_fails_cleanly_without_overwriting(self):
        for m in self.prop['mentions']:
            m.update(visibility='reference', state_id=None, note='仅提及，不展示。')
        del self.prop['states']
        with tempfile.TemporaryDirectory(prefix='prop-missing-states-') as temp:
            root = Path(temp)
            script, assets = root / 'scenes.json', root / 'assets.json'
            script.write_text(json.dumps(self.data, ensure_ascii=False), encoding='utf-8')
            assets.write_text(json.dumps(self.assets, ensure_ascii=False), encoding='utf-8')
            before = assets.read_bytes()
            result = subprocess.run([sys.executable, str(Path(__file__).with_name('build_xlsx.py')),
                                     str(script), str(assets), '--sync-prop-scenes'], capture_output=True, text=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn('缺少非空 states', result.stdout)
            self.assertNotIn('Traceback', result.stderr)
            self.assertEqual(assets.read_bytes(), before)

    def test_local_phrase_can_cover_confirmed_name_substring(self):
        mention = self.prop['mentions'][0]
        text = '凯尔拔出这把屠龙匕首，刀柄刻有三道血焰纹。'
        self.data['episodes'][0]['scenes'][0]['lines'][0]['text'] = text
        mention.update(raw_name='这把屠龙匕首', start_offset=4, end_offset=10,
                       basis='introduced', identity_reason='原文明确凯尔拔出的唯一实物。')
        mention['source']['quote'] = text
        self.prop['mentions'][1]['identity_refs'] = [copy.deepcopy(mention['source'])]
        self.prop['alias_evidence'][0]['evidence_refs'][0] = copy.deepcopy(mention['source'])
        self.sync()

    def test_demonstrative_generic_alias_is_forbidden(self):
        self.prop['aliases'].append('这把匕首')
        self.prop['alias_evidence'].append({'alias': '这把匕首', 'basis': 'continuity', 'reason': '测试',
                                           'evidence_refs': [self.prop['mentions'][0]['source'], self.prop['mentions'][1]['source']]})
        self.rejected('泛称「这把匕首」不能作全局')

    def two_distinct_props(self):
        text = '凯尔左手拿着匕首，右手拿着另一把匕首。'
        self.data['episodes'][0]['scenes'][0]['lines'][0]['text'] = text
        props = []
        for i, low in enumerate([text.index('匕首'), text.rindex('匕首')], 1):
            props.append({'id': f'PRP-00{i}', 'name': f'匕首·凯尔第{i}把', 'owner': '凯尔',
                          'aliases': [], 'alias_evidence': [], 'identity_relations': [],
                          'mentions': [{'id': 'M001', 'source': {'scene_id': '1-1', 'line': 1, 'quote': text},
                                        'raw_name': '匕首', 'start_offset': low, 'end_offset': low + 2,
                                        'basis': 'introduced', 'identity_reason': f'原文明示同时拿着两把不同匕首，区分为第{i}把。',
                                        'timeline': '连续段1', 'visibility': 'visible', 'state_id': 'S01', 'holder': '凯尔'}],
                          'states': [{'id': 'S01', 'name': '常态', 'scenes': []}]})
        self.assets['props'] = props
        return props

    def test_same_type_same_owner_distinct_word_spans_remain_separate(self):
        props = self.two_distinct_props()
        props[0]['identity_relations'] = [{'other_id': props[1]['id'], 'decision': 'distinct',
                                          'reason': '原文明示同时分别持有两把。',
                                          'evidence_refs': [copy.deepcopy(props[0]['mentions'][0]['source'])]}]
        self.sync()
        self.assertEqual(len(self.assets['props']), 2)
        self.assertEqual([p['states'][0]['scenes'] for p in props], [['1-1'], ['1-1']])

    def test_distinct_props_may_share_source_name_not_global_identity(self):
        props = self.two_distinct_props()
        for prop in props:
            prop['source_name'] = '匕首'
        self.sync()
        self.assertEqual(len(self.assets['props']), 2)

    def test_same_word_span_cannot_be_assigned_to_two_physical_props(self):
        props = self.two_distinct_props()
        props[1]['mentions'][0]['start_offset'] = props[0]['mentions'][0]['start_offset']
        props[1]['mentions'][0]['end_offset'] = props[0]['mentions'][0]['end_offset']
        self.rejected('重复认领同一原文道具称呼')

    def test_distinct_relation_requires_original_evidence(self):
        props = self.two_distinct_props()
        props[0]['identity_relations'] = [{'other_id': props[1]['id'], 'decision': 'distinct',
                                          'reason': '分别持有', 'evidence_refs': []}]
        self.rejected('缺少结构化原文依据')

    def test_uncertain_relation_requires_visible_asset_note(self):
        props = self.two_distinct_props()
        props[0]['identity_relations'] = [{'other_id': props[1]['id'], 'decision': 'uncertain',
                                          'reason': '身份需要继续核对', 'evidence_refs': []}]
        self.rejected('note 中标明待确认')
        props[0]['note'] = '待确认两件是否在后文交换。'
        self.sync()

    def test_explicit_alias_statement_with_two_literal_names_passes(self):
        text = '凯尔说明屠龙匕首又称龙骨匕首。'
        scene = self.data['episodes'][0]['scenes'][4]
        scene['lines'][0]['text'] = text
        mention = self.prop['mentions'][-1]
        mention['source']['quote'] = text
        mention.update(start_offset=text.index('屠龙匕首'), end_offset=text.index('屠龙匕首') + 4)
        other = copy.deepcopy(mention)
        other.update(id='M006', raw_name='龙骨匕首', start_offset=text.index('龙骨匕首'), end_offset=text.index('龙骨匕首') + 4)
        self.prop['mentions'].append(other)
        self.prop['alias_evidence'][0].update(basis='explicit', evidence_refs=[copy.deepcopy(mention['source'])])
        self.sync()

    def test_substring_is_not_explicit_alias_identity_proof(self):
        self.prop['alias_evidence'][0].update(basis='explicit', evidence_refs=[self.prop['mentions'][2]['source']])
        self.rejected('显式异称说明须同时包含两个称呼')

    def test_sync_is_idempotent(self):
        self.sync()
        before = copy.deepcopy(self.assets)
        self.sync()
        self.assertEqual(self.assets, before)

    def test_readonly_tools_preserve_assets_and_full_context(self):
        extra = '未命中名称的相邻动作。' + '完整上下文' * 80
        self.data['episodes'][0]['scenes'][0]['lines'].append({'type': 'action', 'text': extra})
        with tempfile.TemporaryDirectory(prefix='prop-readonly-') as temp:
            root = Path(temp)
            script, assets = root / 'scenes.json', root / 'assets.json'
            script.write_text(json.dumps(self.data, ensure_ascii=False), encoding='utf-8')
            assets.write_text(json.dumps(self.assets, ensure_ascii=False), encoding='utf-8')
            before = assets.read_bytes()
            tool = str(Path(__file__).with_name('show_prop.py'))
            result = subprocess.run([sys.executable, tool, str(script), str(assets), 'PRP-001'], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertIn(extra, result.stdout)
            result = subprocess.run([sys.executable, tool, str(script), str(assets), 'PRP-001', '--json'], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertTrue(any(hit['raw_name'] == '它' for hit in json.loads(result.stdout)))
            result = subprocess.run([sys.executable, str(Path(__file__).with_name('enrich_descriptions.py')),
                                     str(script), str(assets)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertIn('1-4:L1 [action]「它」: 她将它折断。', result.stdout)
            self.assertEqual(assets.read_bytes(), before)

    def test_exports_canonical_name_alias_evidence_and_local_names(self):
        self.sync()
        wb = Workbook()
        write_asset_sheet(wb, 'props', '道具', '道具名', [('owner', '归属人物')], self.assets, self.data)
        write_breakdown(wb, self.data, self.assets)
        self.assertEqual(wb['道具'].cell(2, 2).value, '屠龙匕首')
        self.assertIn('三道血焰纹', wb['道具'].cell(2, 12).value)
        self.assertIn('原文：龙骨匕首', wb['场次总表'].cell(4, 9).value)
        self.assertIn('原文：它', wb['场次总表'].cell(5, 9).value)
        self.assertEqual(wb['场次总表'].cell(6, 9).value, '')

    def test_cli_sync_and_export(self):
        with tempfile.TemporaryDirectory(prefix='prop-identity-') as temp:
            root = Path(temp)
            script, assets, output = root / 'scenes.json', root / 'assets.json', root / 'assets.xlsx'
            script.write_text(json.dumps(self.data, ensure_ascii=False), encoding='utf-8')
            assets.write_text(json.dumps(self.assets, ensure_ascii=False), encoding='utf-8')
            command = [sys.executable, str(Path(__file__).with_name('build_xlsx.py')), str(script), str(assets)]
            for mode in ['--sync-prop-scenes', '--check', str(output)]:
                result = subprocess.run(command + [mode], capture_output=True, text=True)
                self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            wb = load_workbook(output)
            self.assertEqual(wb.sheetnames, ['人物', '场景', '道具', '场次总表'])
            self.assertEqual(wb['道具'].cell(1, 12).value, '异称与身份依据')
            wb.close()


if __name__ == '__main__':
    unittest.main(verbosity=2)
