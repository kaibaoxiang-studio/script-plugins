#!/usr/bin/env python3
"""Business acceptance for aliases, states, evidence, packages and XLSX sheets."""
from __future__ import annotations

from copy import deepcopy
import json
from pathlib import Path
import sys
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from asset_core import assets_digest, finalize, initial_work, write_json


def block(kind: str, text: str, ref: str, origin: str = "source", speaker: str = "") -> dict:
    item = {"type": kind, "text": text, "source_refs": [ref], "origin": origin,
            "extracted_text": text, "annotations": []}
    if speaker:
        item["speaker"] = speaker
    return item


def synthetic() -> dict:
    scenes = [
        (1, "1-1", "山门", ["阿云", "赤狐", "守卫们"], [
            block("action", "阿云穿蓝衫，手握铜钥匙；赤狐啼叫，守卫们守在山门。", "L3"),
            block("dialogue", "把钥匙给我。", "L4", speaker="阿云")]),
        (1, "1-2", "山门侧廊", ["云儿", "赤狐"], [
            block("action", "母亲称阿云为云儿。阿云换上红袍，赤狐尾巴扫过门槛。", "L7", "generated")]),
        (2, "2-1", "山门", ["阿云", "赤狐"], [
            block("action", "阿云举起铜钥匙，佛怒火莲在半空凝成九瓣，烧断石门锁链。", "L10")])]
    episodes = []
    for epnum in (1, 2):
        items = []
        for number, sid, location, cast, blocks in scenes:
            if number != epnum:
                continue
            items.append({"id": "scn-" + sid, "number": sid, "source_heading": sid + " " + location + " 外 日",
                          "location": location, "interior_exterior": "外", "time": "日",
                          "characters": cast, "source_refs": ["L" + sid.replace("-", "")], "blocks": blocks})
        episodes.append({"id": "ep-" + str(epnum), "number": epnum, "title": "山门",
                         "summary": "", "source_refs": ["E" + str(epnum)], "scenes": items})
    return {"schema_version": "standardized_screenplay_v1", "input_kind": "script",
            "source_manifest": {"filename": "synthetic.txt", "format": "txt"},
            "needs_creative_pass": False, "episodes": episodes,
            "coverage": {"status": "parsed", "unassigned_source_refs": []}}


def ev(work: dict, sid: str, text: str, ref: str, origin: str = "source") -> dict:
    return {"scene_id": sid, "text": text, "source_refs": [ref], "origin": origin}


def main() -> None:
    with TemporaryDirectory(prefix="asset-acceptance-") as directory:
        root = Path(directory)
        source = root / "standardized_screenplay.json"
        write_json(source, synthetic())
        work = initial_work(source)
        assets = work["assets"]
        chars = assets["characters"]
        ayun = next(item for item in chars if item["name"] == "阿云")
        alias = next(item for item in chars if item["name"] == "云儿")
        chars.remove(alias)
        ayun["tier"] = "主角"
        ayun["aliases"] = ["云儿"]
        ayun["alias_evidence"] = [{"name": "云儿", "evidence": [ev(work, "1-2", "母亲称阿云为云儿。", "L7", "generated")]}]
        ayun["evidence"].extend(alias["evidence"])
        ayun["states"][0]["scenes"].remove("2-1")
        ayun["states"].append({"id": "S02", "name": "红袍", "scenes": ["1-2", "2-1"],
                               "script_facts": [{"text": "红袍", "evidence": [ev(work, "1-2", "阿云换上红袍", "L7", "generated")]}],
                               "evidence": [ev(work, "1-2", "阿云换上红袍", "L7", "generated")]})
        fox = next(item for item in chars if item["name"] == "赤狐")
        fox["asset_type"] = "主要动物"
        fox["tier"] = "重要角色"
        fox["script_facts"] = [{"text": "赤狐尾巴", "evidence": [ev(work, "1-2", "赤狐尾巴扫过门槛。", "L7", "generated")]}]
        guards = next(item for item in chars if item["name"] == "守卫们")
        guards["asset_type"] = "群演组"
        guards["tier"] = "群演"
        for loc in assets["locations"]:
            if loc["name"] == "山门侧廊":
                aside = loc
                break
        assets["locations"].remove(aside)
        mainloc = assets["locations"][0]
        mainloc["aliases"] = ["山门侧廊"]
        mainloc["alias_evidence"] = [{"name": "山门侧廊", "evidence": aside["evidence"]}]
        mainloc["evidence"].extend(aside["evidence"])
        mainloc["subscenes"].extend(aside["subscenes"])
        mainloc["states"][0]["scenes"].append("1-2")
        action1 = ev(work, "1-1", "阿云穿蓝衫，手握铜钥匙；赤狐啼叫，守卫们守在山门。", "L3")
        action3 = ev(work, "2-1", "阿云举起铜钥匙，佛怒火莲在半空凝成九瓣，烧断石门锁链。", "L10")
        assets["props"] = [{"id": "PRP-001", "name": "铜钥匙", "aliases": [], "alias_evidence": [],
                            "script_facts": [], "evidence": [action1, action3],
                            "states": [{"id": "S01", "name": "常态", "scenes": ["1-1", "2-1"],
                                        "script_facts": [], "evidence": [action1]}]}]
        assets["effects"] = [{"id": "VFX-001", "name": "佛怒火莲", "aliases": [], "alias_evidence": [],
                              "script_facts": [{"text": "九瓣", "evidence": [action3]}], "evidence": [action3],
                              "states": [{"id": "S01", "name": "九瓣火莲", "scenes": ["2-1"],
                                          "script_facts": [], "evidence": [action3]}]}]
        work["findings"] = [{"code": "possible_prop_alias", "message": "石门锁链是否需单独制作待确认",
                             "evidence": [action3]}]
        # Authored review fixture: these conclusions are not generated by production code.
        scene_notes = {
            "1-1": (["CHR-001", "CHR-002", "CHR-003"], ["PRP-001"], [],
                    "阿云、赤狐及守卫实际出镜；蓝衫归人物状态。", "手握铜钥匙，有独立使用依据。", "仅人物和实物动作，无独立特效。"),
            "1-2": (["CHR-001", "CHR-002"], [], [],
                    "云儿有明确别名证据归阿云；红袍是新造型状态。", "只描述换装与尾巴动作，门槛为固定陈设。", "赤狐尾巴为动物特征，不立特效。"),
            "2-1": (["CHR-001", "CHR-002"], ["PRP-001"], ["VFX-001"],
                    "阿云与赤狐按本场人物表出镜，阿云红袍状态延续。", "同一铜钥匙再次举起；石门锁链留作制作疑点。", "九瓣火莲有具体外形和因果作用，独立登记。")}
        reviews = []
        for scene in work["scenes"]:
            sid = scene["scene_id"]
            cids, pids, vids, cn, pn, vn = scene_notes[sid]
            blocks = []
            for index, item in enumerate(scene["blocks"]):
                text = (item.get("speaker", "") + "：" if item["type"] == "dialogue" else "") + item["text"]
                linked = (["CHR-001", "PRP-001"] if item["type"] == "dialogue" else cids + pids + vids)
                blocks.append({"block_index": index, "source_refs": item["source_refs"], "origin": item["origin"],
                               "text": text, "asset_ids": linked, "conclusion": cn + pn + vn})
            reviews.append({"scene_id": sid, "blocks": blocks, "categories": {
                "characters": {"asset_ids": cids, "conclusion": cn},
                "locations": {"asset_ids": ["SCN-001"], "conclusion": "同一山门主空间，侧廊归其子场景；局部锁链破损不拆整个场景。"},
                "props": {"asset_ids": pids, "conclusion": pn},
                "effects": {"asset_ids": vids, "conclusion": vn}}, "exclusions": []})
        work["semantic_review"] = {"schema_version": "asset_semantic_review_v1",
                                   "source_sha256": work["source"]["sha256"],
                                   "assets_sha256": assets_digest(assets), "scenes": reviews}
        work_file = root / "work.json"
        write_json(work_file, work)
        result = finalize(work_file, root / "out")
        assert result == {"characters": 3, "locations": 1, "props": 1, "effects": 1,
                          "scenes": 3, "findings": 1}, result
        data = json.loads((root / "out/assets.json").read_text(encoding="utf-8"))
        assert len(data["characters"][0]["requirements"]) >= 8
        assert any(item["type"] == "声音" for item in data["characters"][0]["requirements"])
        assert any(item["type"] == "声音" for item in data["characters"][1]["requirements"])
        assert not any(item["type"] == "声音" for item in data["characters"][2]["requirements"])
        assert data["characters"][0]["states"][1]["evidence"][0]["origin"] == "generated"
        wb = load_workbook(root / "out/screenplay-asset-table.xlsx", read_only=True)
        assert wb.sheetnames == ["人物", "场景", "道具", "场次总表", "特效"], wb.sheetnames
        assert wb["场次总表"].max_row == 4
        assert "待处理" not in [cell.value for cell in wb["场次总表"][1]]
        rows = list(wb["场次总表"].values)
        assert "阿云·红袍" in rows[2][7] and "赤狐·常态" in rows[2][7]
        assert "铜钥匙·常态" in rows[3][8] and "佛怒火莲·九瓣火莲" in rows[3][9]
        assert wb["人物"][2][1].value == "阿云" and wb["道具"][2][1].value == "铜钥匙"
        wb.close()
        bad = deepcopy(work)
        bad["assets"]["props"][0]["evidence"][0]["source_refs"] = ["L999"]
        write_json(work_file, bad)
        try:
            finalize(work_file, root / "bad")
        except ValueError as exc:
            assert "来源位置无效" in str(exc)
        else:
            raise AssertionError("bad source reference was accepted")
        unsupported = deepcopy(work)
        unsupported["assets"]["characters"][0]["script_facts"] = [
            {"text": "银白长发", "evidence": [ev(work, "1-1", "阿云穿蓝衫", "L3")]}]
        write_json(work_file, unsupported)
        try:
            finalize(work_file, root / "unsupported")
        except ValueError as exc:
            assert "事实不是证据原文片段" in str(exc)
        else:
            raise AssertionError("unsupported visual detail was accepted")
        print(json.dumps({"status": "passed", **result, "sheets": wb.sheetnames}, ensure_ascii=False))


if __name__ == "__main__":
    main()
