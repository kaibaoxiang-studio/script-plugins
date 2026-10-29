"""只读验收已保存 Excel：从当前源数据重建期望契约，非从导出计划自证。"""
import argparse
import json
from zipfile import BadZipFile
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from export_contract import build_plan, STYLE, text_lines


def audit_workbook(workbook, plan):
    errors = []
    if workbook.sheetnames != [s['name'] for s in plan['sheets']]:
        errors.append('工作表名称/顺序与标准模板不一致')
    for spec in plan['sheets']:
        if spec['name'] not in workbook.sheetnames: continue
        sheet = workbook[spec['name']]
        label = sheet.title
        expected = [spec['headers']] + spec['rows']
        if (sheet.max_row, sheet.max_column) != (len(expected), len(spec['headers'])):
            errors.append(f'{label} 行列数量错误')
        merges = {}
        merge_ranges = set()
        for group in spec['groups']:
            if group['end'] <= group['start']: continue
            for col in group['merge_columns']:
                name = get_column_letter(col)
                merge_ranges.add(f'{name}{group["start"]}:{name}{group["end"]}')
                for row in range(group['start'], group['end'] + 1):
                    merges[row, col] = group
        if {str(r) for r in sheet.merged_cells.ranges} != merge_ranges:
            errors.append(f'{label} 合并范围错误')
        if sheet.freeze_panes != spec['freeze']: errors.append(f'{label} 冻结位置错误')
        if sheet.auto_filter.ref != spec['filter']: errors.append(f'{label} 缺少完整筛选范围')
        if sheet.sheet_view.showGridLines is not False: errors.append(f'{label} 网格线设置错误')
        for col, width in enumerate(spec['widths'], 1):
            actual = sheet.column_dimensions[get_column_letter(col)].width
            if abs(actual - width) > .05: errors.append(f'{label} 第{col}列宽偏离模板')
        for row, values in enumerate(expected, 1):
            wanted_height = STYLE['header_height'] if row == 1 else spec['row_heights'][row - 2]
            actual_height = sheet.row_dimensions[row].height or 15
            if abs(actual_height - wanted_height) > .1: errors.append(f'{label} 第{row}行高偏离模板')
            group_index = next((i for i, group in enumerate(spec['groups']) if group['start'] <= row <= group['end']), 0)
            wanted_fill = STYLE['header_fill'] if row == 1 else STYLE['group_fills'][group_index % 2]
            for col, wanted in enumerate(values, 1):
                cell = sheet.cell(row, col)
                if (cell.value if cell.value is not None else '') != (wanted if wanted is not None else ''):
                    errors.append(f'{label}!{cell.coordinate} 内容遗漏/排序错误/与数据不一致')
                group = merges.get((row, col))
                if group and row > group['start']: continue
                fill = cell.fill.fgColor
                if fill.type != 'rgb' or str(fill.rgb)[-6:].upper() != wanted_fill:
                    errors.append(f'{label}!{cell.coordinate} 分组底色错误')
                if cell.font.name != STYLE['font'] or cell.font.sz != STYLE['font_size']:
                    errors.append(f'{label}!{cell.coordinate} 字体与模板不一致')
                is_text = col in spec['text_columns']
                vertical = 'center'
                horizontal = 'center'
                if cell.alignment.vertical != vertical or cell.alignment.horizontal != horizontal or not cell.alignment.wrap_text:
                    errors.append(f'{label}!{cell.coordinate} 对齐/换行错误')
                for side in ['left', 'right', 'top', 'bottom']:
                    edge = getattr(cell.border, side)
                    if edge.style != 'thin' or edge.color is None or edge.color.type != 'rgb' or str(edge.color.rgb)[-6:].upper() != STYLE['border']:
                        errors.append(f'{label}!{cell.coordinate} 边框错误'); break
                if row > 1 and wanted:
                    available = sum(sheet.row_dimensions[i].height or 15 for i in range(group['start'], group['end'] + 1)) if group else actual_height
                    needed = text_lines(wanted, spec['widths'][col - 1]) * 15 + 8
                    if needed > available + 1: errors.append(f'{label}!{cell.coordinate} 可能文字截断')
    return errors


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('scenes_file'); parser.add_argument('assets_file'); parser.add_argument('workbook')
    parser.add_argument('--warning-approvals')
    parser.add_argument('--review-file')
    args = parser.parse_args()
    from build_xlsx import validate
    from warning_approvals import apply_approvals
    try:
        with open(args.scenes_file, encoding='utf-8') as f: data = json.load(f)
        with open(args.assets_file, encoding='utf-8') as f: assets = json.load(f)
        document = None
        if args.warning_approvals:
            with open(args.warning_approvals, encoding='utf-8') as f: document = json.load(f)
        errors, warnings = validate(data, assets)
        from pathlib import Path
        from execution_review import load_review, validate_review, source_errors
        errors.extend(source_errors(data))
        review_path=Path(args.review_file or Path(args.scenes_file).with_name('review.json'))
        review=load_review(review_path)
        from audit_state import team_mode_errors
        errors.extend(team_mode_errors(data,review,[review_path.parent,Path(args.scenes_file).parent]))
        errors.extend(validate_review(data, assets, review))
        remaining, accepted, approval_errors = apply_approvals(data, assets, warnings, document)
        errors.extend(approval_errors)
        errors.extend('未确认警告：' + message for message in remaining)
        if not errors:
            workbook=load_workbook(args.workbook)
            try:
                errors = audit_workbook(workbook, build_plan(data, assets, accepted))
            finally:
                workbook.close()
    except (OSError, ValueError, KeyError, BadZipFile) as error:
        parser.exit(1, f'无法完成验收：{error}\n')
    for error in errors: print('错误:', error)
    print(f'Excel 内容与格式验收：{len(errors)} 条错误；仍需正常缩放下人工检查阅读密度和筛选使用。')
    parser.exit(1 if errors else 0)


if __name__ == '__main__': main()
