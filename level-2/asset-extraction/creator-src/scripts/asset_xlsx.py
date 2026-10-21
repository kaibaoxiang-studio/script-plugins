"""Render the four established asset-table sheets and an optional effects sheet."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


HEADER = PatternFill("solid", fgColor="305496")
ALT = PatternFill("solid", fgColor="EEF3FA")
WRAP = Alignment(vertical="top", wrap_text=True)


def words(value) -> str:
    if isinstance(value, list):
        return "、".join(str(item) for item in value)
    return str(value or "")


def refs(evidence: list[dict]) -> str:
    return words(list(dict.fromkeys(ref for ev in evidence for ref in ev.get("source_refs", []))))


def origins(evidence: list[dict]) -> str:
    return words(list(dict.fromkeys(ev.get("origin", "") for ev in evidence if ev.get("origin"))))


def facts(rows: list[dict]) -> str:
    return "\n".join(row["text"] for row in rows)


def styled_sheet(wb: Workbook, title: str, headers: list[str], widths: list[int]):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for idx, cell in enumerate(ws[1], 1):
        cell.fill = HEADER
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = WRAP
        ws.column_dimensions[cell.column_letter].width = widths[idx - 1]
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "C2"
    return ws


def append(ws, values: list, alternating: bool = False) -> None:
    ws.append(values)
    row = ws.max_row
    ws.row_dimensions[row].height = 48
    for cell in ws[row]:
        cell.alignment = WRAP
        if alternating:
            cell.fill = ALT


def state_rows(wb: Workbook, assets: dict, kind: str, sheet: str) -> None:
    if kind == "characters":
        headers = ["资产ID", "人物名", "别名/称谓", "资产类型", "角色层级", "剧本明确特征",
                   "状态ID", "状态名", "状态明确特征", "出现场次", "制作需求", "原文依据", "来源位置", "来源类型"]
        widths = [13, 18, 19, 13, 13, 38, 17, 18, 40, 25, 40, 50, 20, 14]
    elif kind == "locations":
        headers = ["场景ID", "主场景名", "子场景", "剧本明确特征", "状态ID", "状态名",
                   "状态明确特征", "出现场次", "原文依据", "来源位置", "来源类型"]
        widths = [13, 21, 25, 40, 17, 18, 40, 25, 50, 20, 14]
    else:
        headers = ["资产ID", "道具名" if kind == "props" else "特效名", "别名/称谓", "剧本明确特征",
                   "状态ID", "状态名", "状态明确特征", "出现场次", "原文依据", "来源位置", "来源类型"]
        widths = [13, 20, 20, 40, 17, 18, 40, 25, 50, 20, 14]
    ws = styled_sheet(wb, sheet, headers, widths)
    for index, asset in enumerate(assets[kind]):
        for state_index, state in enumerate(asset["states"]):
            ev = state["evidence"]
            common = [asset["id"] if not state_index else "", asset["name"] if not state_index else ""]
            if kind == "characters":
                req = [r["type"] for r in asset.get("requirements", [])
                       if r.get("state_id") in {None, state["id"]}]
                values = common + [words(asset.get("aliases", [])) if not state_index else "",
                                   asset.get("asset_type", "") if not state_index else "",
                                   asset.get("tier", "") if not state_index else "",
                                   facts(asset.get("script_facts", [])) if not state_index else "",
                                   f"{asset['id']}-{state['id']}", state["name"],
                                   facts(state.get("script_facts", [])), words(state["scenes"]),
                                   words(req), "\n".join(item["text"] for item in ev), refs(ev), origins(ev)]
            elif kind == "locations":
                sub = [item["name"] for item in asset["subscenes"]
                       if set(item.get("scenes", [])) & set(state["scenes"])]
                values = common + [words(sub), facts(asset.get("script_facts", [])) if not state_index else "",
                                   f"{asset['id']}-{state['id']}", state["name"],
                                   facts(state.get("script_facts", [])), words(state["scenes"]),
                                   "\n".join(item["text"] for item in ev), refs(ev), origins(ev)]
            else:
                values = common + [words(asset.get("aliases", [])) if not state_index else "",
                                   facts(asset.get("script_facts", [])) if not state_index else "",
                                   f"{asset['id']}-{state['id']}", state["name"],
                                   facts(state.get("script_facts", [])), words(state["scenes"]),
                                   "\n".join(item["text"] for item in ev), refs(ev), origins(ev)]
            append(ws, values, index % 2 == 1)
    ws.auto_filter.ref = ws.dimensions


def at_scene(assets: list[dict], sid: str) -> list[str]:
    return [f"{asset['name']}·{state['name']}"
            for asset in assets for state in asset["states"] if sid in state["scenes"]]


def scene_summary(wb: Workbook, data: dict) -> None:
    has_effects = bool(data["effects"])
    headers = ["场次", "集", "集标题", "场景名(剧本)", "内/外", "时间",
               "场景资产·状态", "人物·状态", "道具·状态"]
    widths = [9, 7, 22, 25, 9, 9, 35, 48, 42]
    if has_effects:
        headers.append("特效·状态")
        widths.append(35)
    headers.append("来源位置")
    widths.append(19)
    ws = styled_sheet(wb, "场次总表", headers, widths)
    for index, scene in enumerate(data["scenes"]):
        sid = scene["scene_id"]
        values = [sid, scene["episode"], scene["episode_title"], scene["location"],
                  scene["interior_exterior"], scene["time"],
                  "\n".join(at_scene(data["locations"], sid)),
                  "\n".join(at_scene(data["characters"], sid)),
                  "\n".join(at_scene(data["props"], sid))]
        if has_effects:
            values.append("\n".join(at_scene(data["effects"], sid)))
        values.append(words(scene["source_refs"]))
        append(ws, values, index % 2 == 1)
    ws.auto_filter.ref = ws.dimensions


def render_workbook(data: dict, target: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    state_rows(wb, data, "characters", "人物")
    state_rows(wb, data, "locations", "场景")
    state_rows(wb, data, "props", "道具")
    scene_summary(wb, data)
    if data["effects"]:
        state_rows(wb, data, "effects", "特效")
    wb.save(target)
