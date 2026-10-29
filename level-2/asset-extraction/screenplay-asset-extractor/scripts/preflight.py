#!/usr/bin/env python3
"""Preflight checks for presentation limits and cross-table references."""
import json
from pathlib import Path
from openpyxl import load_workbook
from export_contract import STYLE, text_lines, WIDTHS

def run(work):
    work=Path(work); errors=[]; warnings=[]
    draft=work/'asset-table.review.xlsx'
    if draft.exists():
        wb=load_workbook(draft, read_only=False, data_only=False)
        for ws in wb.worksheets:
            for row in range(2,ws.max_row+1):
                h=ws.row_dimensions[row].height or 15
                if h > STYLE['maximum_height']: errors.append(f'{ws.title}!{row} 行高 {h:.1f} 超过 {STYLE["maximum_height"]}')
                if h > 360: warnings.append(f'{ws.title}!{row} 行高 {h:.1f} 接近 Excel 上限')
        wb.close()
    scenes=json.loads((work/'scenes.json').read_text())
    assets=json.loads((work/'assets.json').read_text())
    scene_ids=[s['scene_id'] for e in scenes['episodes'] for s in e['scenes']]
    if len(scene_ids)!=len(set(scene_ids)): errors.append('原文存在重复场次编号')
    review=json.loads((work/'review.json').read_text())
    reviewed={x.get('scene_id') for x in review.get('scenes',[])}
    missing=[x for x in scene_ids if x not in reviewed]
    if missing: warnings.append('审核台账缺少场次：'+ '、'.join(missing[:20]))
    valid_scenes=set(scene_ids)
    asset_ids={k:{a.get('id') for a in assets.get(k,[])} for k in ('characters','scenes','props')}
    for kind in ('characters','scenes','props'):
        for asset in assets.get(kind,[]):
            for state in asset.get('states',[]):
                for sid in state.get('scenes',[]):
                    if sid not in valid_scenes: errors.append(f'{kind} {asset.get("id")} 状态 {state.get("id")} 引用不存在场次 {sid}')
            for app in asset.get('appearances',[]):
                sid=app.get('scene_id')
                if sid and sid not in valid_scenes: errors.append(f'{kind} {asset.get("id")} 出场引用不存在场次 {sid}')
    for scene in review.get('scenes',[]):
        for line in scene.get('lines',[]):
            for aid in line.get('asset_ids',[]):
                if not any(aid in ids for ids in asset_ids.values()): errors.append(f'{scene.get("scene_id")} 原文行引用不存在资产 {aid}')
    return {'ok':not errors,'errors':errors,'warnings':warnings,'scene_count':len(scene_ids),'reviewed_scene_count':len(reviewed)}

