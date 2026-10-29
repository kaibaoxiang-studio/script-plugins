#!/usr/bin/env python3
"""Generate representative PNG previews for the four-sheet review gate."""
import hashlib, json, shutil, subprocess, tempfile
from pathlib import Path
from openpyxl import load_workbook

SHEETS = ('人物', '场景', '道具', '场次总表')

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def generate(work):
    work = Path(work).resolve(); draft = work / 'asset-table.review.xlsx'
    if not draft.is_file(): raise ValueError('缺少 asset-table.review.xlsx；先 workflow.py export')
    evidence = work / 'visual_evidence'; evidence.mkdir(exist_ok=True)
    missing=[name for name in ('soffice','pdftoppm') if shutil.which(name) is None]
    if missing: raise ValueError('视觉预览缺少依赖：' + '、'.join(missing) + '；请安装 LibreOffice 和 Poppler')
    temp = Path(tempfile.mkdtemp(prefix='asset-preview-'))
    try:
        wb = load_workbook(draft)
        for sheet in SHEETS:
            one = load_workbook(draft)
            for name in list(one.sheetnames):
                if name != sheet: del one[name]
            ws = one[sheet]
            ws.page_setup.orientation = 'landscape'
            ws.page_setup.paperSize = ws.PAPERSIZE_A3
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
            ws.print_title_rows = '1:1'
            out = temp / f'{sheet}.xlsx'; one.save(out); one.close()
            subprocess.run(['soffice','--headless','--convert-to','pdf','--outdir',str(temp),str(out)], check=True, capture_output=True)
            pdf = temp / f'{sheet}.pdf'
            import subprocess as sp
            pages = int(sp.check_output(['pdfinfo',str(pdf)], text=True).split('Pages:')[1].splitlines()[0].strip()) if shutil.which('pdfinfo') else 1
            picks = sorted(set([1, max(1, (pages+1)//2), pages]))
            files=[]
            for page in picks:
                prefix=temp/f'{sheet}-page-{page}'
                sp.run(['pdftoppm','-f',str(page),'-l',str(page),'-png','-r','90',str(pdf),str(prefix)], check=True, capture_output=True)
                candidates=sorted(temp.glob(f'{sheet}-page-{page}-*.png'))
                if not candidates: raise ValueError(f'无法找到 {sheet} 第 {page} 页预览')
                produced=candidates[0]
                target=evidence/f'{sheet}-自动预览-{page}.png'; shutil.copyfile(produced,target); files.append(str(target))
            png=files[0]
        wb.close()
    finally:
        shutil.rmtree(temp, ignore_errors=True)
    rows=[]
    for sheet in SHEETS:
        p=evidence/f'{sheet}-自动预览.png'
        files=sorted(str(x) for x in evidence.glob(f'{sheet}-自动预览-*.png'))
        rows.append({'sheet':sheet,'conclusion':'','evidence_file':files[0],
                     'evidence_files':files,'evidence_sha256':sha(Path(files[0])),
                     'evidence_sha256s':[sha(Path(x)) for x in files]})
    return rows

if __name__ == '__main__':
    import argparse
    ap=argparse.ArgumentParser(); ap.add_argument('work'); args=ap.parse_args()
    print(json.dumps({'visual':generate(args.work)},ensure_ascii=False,indent=2))
