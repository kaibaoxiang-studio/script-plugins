#!/usr/bin/env python3
"""阶段入口及分工提交、分批校验、独立审核。不会自动写已读或已复核结论。"""
import argparse
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from datetime import datetime, timezone
from zipfile import BadZipFile
from docx.opc.exceptions import PackageNotFoundError
from openpyxl import load_workbook
from execution_review import (VERSION, digest, file_digest, asset_digest, skeleton, load_review,
                              validate_review, source_errors, concrete)

SCRIPTS = Path(__file__).resolve().parent
DRAFT = 'asset-table.review.xlsx'


def read(path):
    value = json.loads(Path(path).read_text(encoding='utf-8'))
    if not isinstance(value, dict):
        raise ValueError(f'{Path(path).name} 顶层必须为对象')
    return value


def write(path, value):
    path = Path(path)
    descriptor, temporary = tempfile.mkstemp(prefix='.asset-state-', suffix='.json', dir=path.parent)
    try:
        with os.fdopen(descriptor, 'w', encoding='utf-8') as stream:
            json.dump(value, stream, ensure_ascii=False, indent=2)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def require_no_errors(errors):
    if errors:
        raise ValueError('\n'.join(errors))


def read_inputs(work):
    data, assets, review = read(work / 'scenes.json'), read(work / 'assets.json'), load_review(work / 'review.json')
    if (work / 'team.json').exists():
        from team_workflow import _identity
        team = read(work / 'team.json')
        _identity(team, data)
        if team.get('status') != 'merged_pending_bind':
            raise ValueError('分工片段尚未全部合并；先完成 submit 和 merge')
        if not isinstance(review.get('independent_review'), dict) or review['independent_review'].get('required') is not True:
            raise ValueError('分工模式的独立审核配置丢失；不能跳过审核直接绑定')
    from build_xlsx import validate
    errors, warnings = validate(data, assets)
    require_no_errors(errors)  # 类型/业务错误先返回，不把损坏输入传给后续校验器。
    from warning_approvals import apply_approvals
    approval_file = work / 'warning_approvals.json'
    document = read(approval_file) if approval_file.exists() else None
    remaining, accepted, approval_errors = apply_approvals(data, assets, warnings, document)
    errors.extend(approval_errors)
    errors.extend('未确认警告：' + item for item in remaining)
    errors.extend(source_errors(data))
    require_no_errors(errors)
    return data, assets, review, accepted


def prepare(source, work):
    from parse_script import parse, quality_errors
    if work.exists() and any(work.iterdir()):
        raise ValueError('工作目录已有文件；请为重新解析选择新目录，保留旧记录以便回查')
    if work.resolve() == source.resolve() or source.suffix.lower() not in ('.docx', '.md', '.markdown'):
        raise ValueError('输入必须是标准 DOCX 或 Markdown，工作目录须独立')
    data = parse(source)
    require_no_errors(quality_errors(data))
    work.mkdir(parents=True, exist_ok=True)
    write(work / 'scenes.json', data)
    write(work / 'review.json', skeleton(data))
    write(work / 'progress.json', {'version': VERSION, 'stage': 'parsed',
                                  'source_sha256': digest(data), 'next': '逐场填写资产及 review.json 空白判读'})
    print('解析通过；已建立空白逐场台账。下一步通读并创建 assets.json，不是已完成提取。')


def check(work):
    """Run the complete non-visual gate with one stable command."""
    report = {'stage': 'check', 'errors': [], 'warnings': [], 'workbook': None, 'visual': None}
    try:
        data, assets, review, accepted = read_inputs(work)
        report['warnings'] = [f'已确认警告：{len(accepted)}'] if accepted else []
        report['errors'].extend(validate_review(data, assets, review))
        from audit_state import audit_errors
        report['errors'].extend(audit_errors(data, assets, review))
        report['visual'] = visual_errors(work, review)
        from preflight import run as run_preflight
        report['preflight'] = run_preflight(work)
        report['errors'].extend(report['preflight'].get('errors', []))
        if (work / DRAFT).is_file():
            from export_contract import build_plan
            from check_workbook import audit_workbook
            wb = load_workbook(work / DRAFT)
            try:
                report['workbook'] = audit_workbook(wb, build_plan(data, assets, accepted))
            finally:
                wb.close()
        else:
            report['workbook'] = ['缺少 asset-table.review.xlsx；先 workflow.py export']
    except (OSError, ValueError, KeyError, TypeError, BadZipFile) as error:
        report['errors'].append(str(error))
    report['ok'] = not any(report[key] for key in ('errors', 'workbook', 'visual'))
    return report


def visual_review(work):
    from visual_review import generate
    review = read(work / 'review.json')
    previous = {x.get('sheet'): x for x in (review.get('delivery_review') or {}).get('visual', []) if isinstance(x, dict)}
    rows = generate(work)
    for row in rows:
        old = previous.get(row['sheet'], {})
        if old.get('conclusion'):
            row['conclusion'] = old['conclusion']
    delivery = review.get('delivery_review') or {}
    delivery['visual'] = rows
    review['delivery_review'] = delivery
    write(work / 'review.json', review)
    print(json.dumps({'generated': [row['evidence_file'] for row in rows],
                      'next': '实际查看四张 PNG 后填写 conclusion，再运行 workflow.py check'}, ensure_ascii=False, indent=2))


def input_fingerprints(work):
    names = ('scenes.json', 'assets.json', 'review.json', 'warning_approvals.json', 'team.json')
    result = {work / name: file_digest(work/name) if (work/name).exists() else None for name in names}
    data = read(work/'scenes.json')
    source = Path(data['source']['path'])
    result[source] = file_digest(source)
    return result


def check_unchanged(expected):
    if any((file_digest(path) if path.exists() else None) != wanted for path, wanted in expected.items()):
        raise ValueError('校验期间输入发生并发修改；本次不绑定或交付，重新读取当前数据')


def reusable_delivery(work, data, assets, review, accepted):
    from export_contract import build_plan
    from check_workbook import audit_workbook
    delivery = review.get('delivery_review')
    if not isinstance(delivery, dict) or not (work / DRAFT).is_file():
        return False
    if delivery.get('source_sha256') != digest(data) or delivery.get('plan_sha256') != digest(build_plan(data, assets, accepted)):
        return False
    if delivery.get('workbook_sha256') != file_digest(work / DRAFT):
        return False
    try:
        wb = load_workbook(work / DRAFT)
        try:
            return not audit_workbook(wb, build_plan(data, assets, accepted))
        finally:
            wb.close()
    except (OSError, ValueError, KeyError, BadZipFile):
        return False


def bind(work):
    expected = input_fingerprints(work)
    data, assets, review, accepted = read_inputs(work)
    from audit_state import audit_errors, scope_snapshots
    require_no_errors(validate_review(data, assets, review, bound=False) + audit_errors(data, assets, review))
    keep_visual = reusable_delivery(work, data, assets, review, accepted)
    review['source_sha256'], review['assets_sha256'] = digest(data), asset_digest(assets)
    review['scope_snapshots'] = scope_snapshots(data, assets, review)
    if not keep_visual:
        review['delivery_review'] = {}
    check_unchanged(expected)
    write(work / 'review.json', review)
    write(work / 'progress.json', {'version': VERSION, 'stage': 'reviewed',
                                  'source_sha256': digest(data), 'assets_sha256': asset_digest(assets),
                                  'next': 'workflow.py export 生成并验收复核副本'})
    print('已绑定当前判读与原文/资产快照；未代替语义判断。')
    if keep_visual:
        print('展示计划及工作簿未变化，保留现有视觉记录。')


def export(work):
    expected = input_fingerprints(work)
    data, assets, review, accepted = read_inputs(work)
    require_no_errors(validate_review(data, assets, review))
    from audit_state import audit_errors
    require_no_errors(audit_errors(data, assets, review))
    if reusable_delivery(work, data, assets, review, accepted):
        check_unchanged(expected)
        print('当前复核副本内容与展示计划一致，保留工作簿和已有视觉记录。')
        return
    command = [sys.executable, str(SCRIPTS / 'build_xlsx.py'), str(work / 'scenes.json'),
               str(work / 'assets.json'), str(work / DRAFT), '--review-file', str(work / 'review.json')]
    if (work / 'warning_approvals.json').exists():
        command += ['--warning-approvals', str(work / 'warning_approvals.json')]
    result = subprocess.run(command, capture_output=True, text=True)
    if result.returncode:
        raise ValueError(result.stdout + result.stderr)
    from export_contract import build_plan
    review['delivery_review'] = {'workbook_sha256': file_digest(work / DRAFT),
                                'source_sha256': digest(data), 'plan_sha256': digest(build_plan(data, assets, accepted)),
                                'visual': [{'sheet': sheet, 'conclusion': '', 'evidence_file': '', 'evidence_sha256': ''}
                                           for sheet in ('人物', '场景', '道具', '场次总表')]}
    check_unchanged(expected)
    write(work / 'review.json', review)
    write(work / 'progress.json', {'version': VERSION, 'stage': 'workbook_audited',
                                  'draft_sha256': file_digest(work / DRAFT),
                                  'next': '实际查看四表、筛选并记录预览证据，再 workflow.py finalize'})
    print(result.stdout.strip())
    print('已建立待视觉复核记录；填写四张表的实际结论和本地预览图片证据。')


def visual_errors(work, review):
    errors = []
    delivery = review.get('delivery_review')
    draft = work / DRAFT
    if not isinstance(delivery, dict) or not draft.is_file():
        return ['缺少已自动验收的复核工作簿，先 workflow.py export']
    if delivery.get('workbook_sha256') != file_digest(draft):
        errors.append('工作簿哈希变化或未绑定，重新导出并视觉复核')
    rows = delivery.get('visual', [])
    seen = set()
    if not isinstance(rows, list):
        return errors + ['视觉复核必须为逐表数组']
    for item in rows:
        if not isinstance(item, dict) or item.get('sheet') not in ('人物', '场景', '道具', '场次总表'):
            errors.append('视觉复核引用不存在的表')
            continue
        sheet = item['sheet']
        if sheet in seen:
            errors.append(f'{sheet} 视觉复核重复')
        seen.add(sheet)
        if not concrete(item.get('conclusion')):
            errors.append(f'{sheet} 缺少阅读密度、分组、截断及实际筛选的复核结论')
        name = item.get('evidence_file')
        if not isinstance(name, str) or not name:
            errors.append(f'{sheet} 缺少已查看的本地预览图片证据')
            continue
        path = Path(name)
        if not path.is_absolute():
            path = work / path
        try:
            payload = path.read_bytes()
            from image_evidence import validate_png
            validate_png(payload)
            if file_digest(path) != item.get('evidence_sha256'):
                errors.append(f'{sheet} 预览图片无效或哈希未绑定')
        except ValueError as error:
            errors.append(f'{sheet} 预览图片无效：{error}')
        except OSError:
            errors.append(f'{sheet} 预览证据文件不存在')
    if seen != {'人物', '场景', '道具', '场次总表'}:
        errors.append('四张表的视觉复核未完成')
    return errors


def finalize(work, output):
    expected = input_fingerprints(work)
    data, assets, review, accepted = read_inputs(work)
    require_no_errors(validate_review(data, assets, review) + visual_errors(work, review))
    from export_contract import build_plan
    from check_workbook import audit_workbook
    draft = load_workbook(work / DRAFT)
    try:
        require_no_errors(audit_workbook(draft, build_plan(data, assets, accepted)))
    finally:
        draft.close()
    output = output.resolve()
    if output.suffix.lower() != '.xlsx' or output == (work / DRAFT).resolve():
        raise ValueError('正式输出须为独立的 .xlsx 文件，不能覆盖复核副本')
    if output == Path(data['source']['path']).resolve() or output in {p.resolve() for p in SCRIPTS.parent.rglob('*') if p.is_file()}:
        raise ValueError('正式输出不能覆盖源剧本或 skill 文件')
    output.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix='.asset-publish-', suffix='.xlsx', dir=output.parent)
    os.close(descriptor)
    try:
        shutil.copyfile(work / DRAFT, temporary)
        if file_digest(temporary) != review['delivery_review']['workbook_sha256']:
            raise ValueError('交付副本与已验收工作簿不一致')
        check_unchanged(expected)
        os.replace(temporary, output)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
    report = {'schema_version': 'asset_acceptance_v1', 'version': VERSION,
              'created_at': datetime.now(timezone.utc).isoformat(),
              'source_docx': data['source'], 'source_sha256': digest(data),
              'assets_sha256': asset_digest(assets), 'assets_file_sha256': file_digest(work / 'assets.json'),
              'review_file_sha256': file_digest(work / 'review.json'),
              'warning_approvals_sha256': file_digest(work / 'warning_approvals.json') if (work / 'warning_approvals.json').exists() else None,
              'output': str(output), 'workbook_sha256': file_digest(output),
              'errors': 0, 'unconfirmed_warnings': 0, 'confirmed_warnings': len(accepted),
              'confirmed_warning_details': accepted,
              'parse_unrecognized': data.get('parse_report', {}).get('unrecognized_count', 0),
              'parse_reviewed': len(review.get('parse_reviews', [])),
              'reviewed_scenes': len(review['scenes']), 'reviewed_candidates': len(review['candidate_reviews']),
              'counts': {kind: {'assets': len(assets[kind]), 'states': sum(len(a.get('states', [])) for a in assets[kind])}
                         for kind in ('characters', 'scenes', 'props')},
              'checks': {'data': 'passed', 'workbook': 'passed', 'visual': 'evidence_recorded',
                         'semantic': ('independent_scope_review_recorded' if review.get('independent_review') else 'review_recorded') + '; correctness still depends on actual reading'}}
    if review.get('independent_review'):
        state=review['independent_review']
        report['independent_audit']={'scenes':len(state['scenes']),'characters':len(state['characters']),
                                    'unresolved_issues':sum(i.get('status')!='resolved' for i in state['issues']),
                                    'reviewers':sorted({v['reviewer_id'] for k in ('scenes','characters') for v in state[k].values()})}
    write(work / 'acceptance.json', report)
    from delivery_report import generate as generate_delivery_report
    (work / '交付报告.md').write_text(generate_delivery_report(work), encoding='utf-8')
    write(work / 'progress.json', {'version': VERSION, 'stage': 'complete', 'acceptance': 'acceptance.json'})
    print('正式交付验收通过：' + str(output))
    print(json.dumps(report['counts'], ensure_ascii=False))


def status(work):
    if not work.exists():
        return {'stage': 'not_started', 'next': 'workflow.py prepare'}
    if not (work / 'scenes.json').is_file():
        return {'stage': 'not_started', 'issues': ['缺少 scenes.json，先重新解析来源']}
    try:
        data = read(work / 'scenes.json')
        errors = source_errors(data)
        if errors:
            return {'stage': 'source_invalid', 'issues': errors[:20]}
        if not (work / 'assets.json').is_file():
            return {'stage': 'extraction_pending', 'next': '通读剧本，创建 assets.json 并填写逐场台账'}
        if not (work / 'review.json').is_file():
            return {'stage': 'review_pending', 'issues': ['缺少 review.json，须回查原文补审核']}
        data, assets, review, accepted = read_inputs(work)
    except (OSError, ValueError, KeyError, TypeError) as error:
        return {'stage': 'data_invalid', 'issues': [str(error)]}
    errors = validate_review(data, assets, review)
    if errors:
        from audit_state import changed_scopes
        return {'stage': 'review_pending', 'issues': errors, 'affected_scopes': changed_scopes(data, assets, review)}
    from audit_state import audit_errors
    errors = audit_errors(data, assets, review)
    if errors:
        return {'stage': 'independent_review_pending', 'issues': errors}
    errors = visual_errors(work, review)
    if errors:
        return {'stage': 'visual_pending', 'issues': errors[:20]}
    from export_contract import build_plan
    from check_workbook import audit_workbook
    try:
        workbook = load_workbook(work / DRAFT)
        try:
            errors = audit_workbook(workbook, build_plan(data, assets, accepted))
        finally:
            workbook.close()
    except (OSError, ValueError, KeyError, BadZipFile) as error:
        return {'stage': 'workbook_invalid', 'issues': [str(error)]}
    if errors:
        return {'stage': 'workbook_invalid', 'issues': errors[:20]}
    receipt = work / 'acceptance.json'
    if not receipt.exists():
        return {'stage': 'ready_to_finalize'}
    value = read(receipt)
    output = Path(value.get('output', ''))
    current = {'source_sha256': digest(data), 'assets_sha256': asset_digest(assets),
               'assets_file_sha256': file_digest(work / 'assets.json'), 'review_file_sha256': file_digest(work / 'review.json'),
               'warning_approvals_sha256': file_digest(work / 'warning_approvals.json') if (work / 'warning_approvals.json').exists() else None,
               'version': VERSION}
    if any(value.get(key) != wanted for key, wanted in current.items()) or not output.is_file() or value.get('workbook_sha256') != file_digest(output):
        return {'stage': 'acceptance_expired', 'issues': ['文件或验收快照变化，重新验收']}
    return {'stage': 'complete', 'output': str(output), 'workbook_sha256': value['workbook_sha256']}


def emit_report(parser, report):
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if report.get('issues') or report.get('errors') or report.get('global_issues'):
        parser.exit(1)


def template_output(parser, report, key, output, work):
    if report.get('issues'):
        emit_report(parser, report)
    value = report[key] if key else report
    if output is None:
        print(json.dumps(value, ensure_ascii=False, indent=2))
        return
    output = output.resolve()
    if output.exists() or output.suffix.lower() != '.json' or output == work / 'scenes.json':
        raise ValueError('模板输出须为新的独立 JSON，不能覆盖已有文件')
    output.parent.mkdir(parents=True, exist_ok=True)
    write(output, value)
    print('已建立待填写模板：' + str(output))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest='command', required=True)
    first = commands.add_parser('prepare'); first.add_argument('source', type=Path); first.add_argument('work', type=Path)
    for name in ('bind', 'export', 'finalize', 'status', 'check', 'visual-review', 'preflight', 'warning-groups', 'warning-draft', 'compare-sources', 'compress-advice', 'impact', 'semantic-audit', 'visual-regression', 'changelog', 'report', 'diagnose-markdown', 'team-init', 'fragment-template', 'submit', 'merge',
                 'check-batch', 'audit-template', 'submit-audit', 'diagnose', 'candidate-groups', 'record-run'):
        command = commands.add_parser(name)
        if name == 'compare-sources':
            command.add_argument('left', type=Path); command.add_argument('right', type=Path)
        elif name == 'warning-draft':
            command.add_argument('work', type=Path)
            command.add_argument('--output', type=Path)
        elif name == 'changelog':
            command.add_argument('work', type=Path); command.add_argument('message')
        elif name == 'diagnose-markdown':
            command.add_argument('source', type=Path)
        else:
            command.add_argument('work', type=Path)
        if name == 'record-run':
            command.add_argument('--kind',choices=('characters','scenes','props'),required=True)
            command.add_argument('--state',choices=('running','failed'),required=True)
            command.add_argument('--agent',default='')
            command.add_argument('--error',default='')
        if name == 'finalize':
            command.add_argument('--output', type=Path, required=True)
        if name in ('fragment-template', 'check-batch'):
            command.add_argument('--kind', choices=('characters', 'scenes', 'props'), required=name=='fragment-template')
        if name in ('fragment-template', 'check-batch', 'audit-template'):
            command.add_argument('--scene', action='append', dest='scenes')
        if name == 'audit-template':
            command.add_argument('--character', action='append', dest='characters')
        if name in ('fragment-template','audit-template'):
            command.add_argument('--output', type=Path)
        if name in ('submit','submit-audit'):
            command.add_argument('--fragment' if name=='submit' else '--report', type=Path, required=True)
    args = parser.parse_args()
    try:
        if hasattr(args, 'work'):
            args.work = args.work.resolve()
        if hasattr(args, 'source'):
            args.source = args.source.resolve()
        if args.command == 'prepare': prepare(args.source.resolve(), args.work)
        elif args.command in ('bind','export','finalize'):
            from team_workflow import _lock
            with _lock(args.work):
                if args.command=='bind': bind(args.work)
                elif args.command=='export': export(args.work)
                else: finalize(args.work,args.output)
        elif args.command == 'status': print(json.dumps(status(args.work), ensure_ascii=False, indent=2))
        elif args.command == 'check':
            report = check(args.work)
            print(json.dumps(report, ensure_ascii=False, indent=2))
            if not report.get('ok'):
                parser.exit(1)
        elif args.command == 'visual-review': visual_review(args.work)
        elif args.command == 'preflight':
            from preflight import run
            emit_report(parser, run(args.work))
        elif args.command == 'warning-groups':
            from warning_groups import group_warnings
            data=read(args.work/'scenes.json'); assets=read(args.work/'assets.json')
            print(json.dumps(group_warnings(data,assets),ensure_ascii=False,indent=2))
        elif args.command == 'warning-draft':
            from warning_groups import approval_draft
            data=read(args.work/'scenes.json'); assets=read(args.work/'assets.json')
            output=args.output or (args.work/'warning_approvals.draft.json')
            write(output, approval_draft(data,assets)); print('已生成批量确认草稿：'+str(output))
        elif args.command == 'compare-sources':
            from source_compare import compare
            print(json.dumps(compare(args.left,args.right),ensure_ascii=False,indent=2))
        elif args.command == 'compress-advice':
            from compress_advice import run
            emit_report(parser, run(args.work))
        elif args.command == 'impact':
            from impact import run
            print(json.dumps(run(args.work),ensure_ascii=False,indent=2))
        elif args.command == 'semantic-audit':
            from semantic_audit import run
            emit_report(parser, run(args.work))
        elif args.command == 'visual-regression':
            from visual_regression import run
            emit_report(parser, run(args.work))
        elif args.command == 'changelog':
            from change_log import append
            print('已写入：'+append(args.work,args.message))
        elif args.command == 'report':
            from delivery_report import generate
            print(generate(args.work))
        elif args.command == 'diagnose-markdown':
            from parse_script import parse, quality_errors
            source = args.source.resolve(); data = parse(source)
            report = {'source': data['source'], 'episodes': len(data['episodes']),
                      'scenes': sum(len(e['scenes']) for e in data['episodes']),
                      'unrecognized_count': data['parse_report']['unrecognized_count'],
                      'quality_errors': quality_errors(data),
                      'unrecognized': data['parse_report'].get('unrecognized', [])}
            print(json.dumps(report, ensure_ascii=False, indent=2))
        elif args.command in ('team-init','fragment-template','submit','merge','check-batch'):
            import team_workflow as team
            if args.command == 'team-init': report=team.init_team(args.work)
            elif args.command == 'fragment-template':
                report=team.fragment_template(args.work,args.kind,args.scenes)
                template_output(parser,report,'fragment',args.output,args.work)
                return
            elif args.command == 'submit': report=team.submit_fragment(args.work,args.fragment)
            elif args.command == 'merge': report=team.merge_fragments(args.work)
            else:
                data=read(args.work/'scenes.json')
                scope=args.scenes or [s['scene_id'] for ep in data['episodes'] for s in ep['scenes']]
                report=team.check_batch(args.work,scope,args.kind)
            emit_report(parser,report)
        elif args.command in ('audit-template','submit-audit'):
            from audit_state import audit_template,submit_audit
            if args.command == 'audit-template':
                data,assets,review=(read(args.work/name) for name in ('scenes.json','assets.json','review.json'))
                require_no_errors(source_errors(data))
                report=audit_template(data,assets,review,args.scenes,args.characters)
                template_output(parser,report,None,args.output,args.work)
            else: emit_report(parser,submit_audit(args.work,args.report))
        elif args.command == 'record-run':
            from team_workflow import record_run
            emit_report(parser,record_run(args.work,args.kind,args.state,args.agent,args.error))
        elif args.command == 'candidate-groups':
            from review_diagnostics import candidate_groups
            data=read(args.work/'scenes.json')
            require_no_errors(source_errors(data))
            print(json.dumps(candidate_groups(data),ensure_ascii=False,indent=2))
        elif args.command == 'diagnose':
            from review_diagnostics import diagnostics
            data,assets,review=(read(args.work/name) for name in ('scenes.json','assets.json','review.json'))
            report=diagnostics(data,assets,review,bound=True)
            from audit_state import changed_scopes
            report['affected_scopes']=changed_scopes(data,assets,review)
            write(args.work/'diagnostics.json',report)
            emit_report(parser,report)
    except (OSError, ValueError, KeyError, TypeError, BadZipFile, PackageNotFoundError) as error:
        parser.exit(1, f'执行未完成：{error}\n')


if __name__ == '__main__':
    main()
