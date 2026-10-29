#!/usr/bin/env python3
"""Report long cells and conservative compression suggestions."""
import json
from pathlib import Path
from openpyxl import load_workbook
from export_contract import STYLE
def run(work):
    p=Path(work)/'asset-table.review.xlsx'; result=[]
    if not p.exists(): return {'ok':False,'issues':['缺少 asset-table.review.xlsx']}
    wb=load_workbook(p,read_only=False,data_only=False)
    for ws in wb.worksheets:
        for row in range(2,ws.max_row+1):
            h=ws.row_dimensions[row].height or 15
            if h>=300:
                vals=[str(ws.cell(row,c).value or '') for c in range(1,ws.max_column+1)]
                long=max(vals,key=len) if vals else ''
                result.append({'sheet':ws.title,'row':row,'height':h,'suggestion':'压缩相邻重复场次/依据，保留完整原文引文；不要删资产或状态。','sample':long[:120]})
    wb.close(); return {'ok':True,'items':result}
