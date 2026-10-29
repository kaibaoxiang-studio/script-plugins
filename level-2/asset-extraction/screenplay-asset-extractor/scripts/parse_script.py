#!/usr/bin/env python3
"""把标准剧本 docx 解析为按场次组织的 JSON。

用法:
    python3 parse_script.py <剧本.docx> <输出scenes.json>

标准剧本格式:
    第 N 集｜标题
    本集概要：……
    N-M 场景名 内/外 时间
    人物：A、B、C
    ▲ 动作描述
    角色（情绪）[VO/OS]：台词
"""
import json
import re
import sys
import hashlib
from pathlib import Path

import docx

RE_EPISODE = re.compile(r'^第\s*(\d+)\s*集(?:\s*[｜|]\s*(.*))?$')
RE_SUMMARY = re.compile(r'^本集概要[：:]\s*(.*)$')
RE_SCENE = re.compile(r'^(\d+)-(\d+)\s+(.+?)\s+(内|外|内外)\s+(\S+)$')
RE_SCENE_ALT = re.compile(r'^(\d+)-(\d+)\s+(日|夜)\s+(内景|外景|内外)\s+(.+)$')
RE_CHARS = re.compile(r'^人物[：:]\s*(.*)$')
RE_DIALOG = re.compile(r'^([^：:▲]{1,15}?)\s*(?:（([^）]*)）)?\s*(?:\[(VO|OS)\])?\s*[：:]\s*(.*)$')


def _markdown_paragraphs(path):
    """Read the supported Markdown screenplay subset without losing source line numbers."""
    rows = []
    for index, raw in enumerate(Path(path).read_text(encoding='utf-8-sig').splitlines(), 1):
        line = raw.strip()
        if not line or line == '---':
            continue
        # Markdown headings and blockquotes are presentation, not screenplay content.
        line = re.sub(r'^#{1,6}\s*', '', line)
        line = re.sub(r'^>\s*', '', line)
        line = re.sub(r'^\*+(.*?)\*+$', r'\1', line)
        line = line.replace('\\-', '-').replace('\\.', '.')
        if line:
            rows.append((index, line))
    return rows, []


def _docx_paragraphs(path):
    doc = docx.Document(path)
    return [(i, p.text.strip()) for i, p in enumerate(doc.paragraphs, 1) if p.text.strip()], doc


def parse(path):
    path = Path(path)
    if path.suffix.lower() in ('.md', '.markdown'):
        paras, doc = _markdown_paragraphs(path)
        tables = []
    else:
        paras, doc = _docx_paragraphs(path)
        tables = doc.tables

    result = {'title': '', 'episodes': [], 'warnings': []}
    ep = None
    scene = None
    unknown = []
    run, maximum_run = 0, 0

    def unrecognized(index, text):
        nonlocal run, maximum_run
        unknown.append({'paragraph': index, 'text': text})
        run += 1
        maximum_run = max(maximum_run, run)

    for index, line in paras:
        m = RE_EPISODE.match(line)
        if m:
            ep = {'episode': int(m.group(1)), 'title': m.group(2), 'summary': '', 'scenes': []}
            result['episodes'].append(ep)
            scene = None
            run = 0
            continue
        if ep is None:
            # 集标题之前的内容视为剧名
            if not result['title']:
                result['title'] = line
                run = 0
            else:
                unrecognized(index, line)
                result['warnings'].append(f'集标题前的未识别行: {line}')
            continue
        m = RE_SUMMARY.match(line)
        if m:
            ep['summary'] = m.group(1)
            run = 0
            continue
        scene_candidate = line[1:].strip() if line.startswith('▲') else line
        scene_candidate = re.sub(r'^\*+(.*?)\*+$', r'\1', scene_candidate)
        scene_candidate = scene_candidate.replace('\\-', '-').replace('\\.', '.')
        m = RE_SCENE.match(scene_candidate)
        if not m:
            alt = RE_SCENE_ALT.match(scene_candidate)
            if alt:
                m = (alt, True)
        if isinstance(m, tuple):
            alt, _ = m
            scene_id = f'{alt.group(1)}-{alt.group(2)}'
            # The alternate form is: scene-id time int/ext location.
            time, int_ext, location = alt.group(3), alt.group(4), alt.group(5)
            scene = {
                'scene_id': scene_id, 'location': location,
                'int_ext': int_ext.replace('内景', '内').replace('外景', '外'),
                'time': time, 'characters': [], 'lines': [],
            }
            ep['scenes'].append(scene); run = 0; continue
        if m:
            scene = {
                'scene_id': f'{m.group(1)}-{m.group(2)}',
                'location': m.group(3),
                'int_ext': m.group(4),
                'time': m.group(5),
                'characters': [],
                'lines': [],
            }
            ep['scenes'].append(scene)
            run = 0
            continue
        if scene is None:
            result['warnings'].append(f'第{ep["episode"]}集场头前的未识别行: {line}')
            unrecognized(index, line)
            continue
        m = RE_CHARS.match(line)
        if m:
            scene['characters'] = [c.strip() for c in re.split(r'[、，,]', m.group(1)) if c.strip()]
            run = 0
            continue
        if line.startswith('▲'):
            scene['lines'].append({'type': 'action', 'text': line.lstrip('▲').strip()})
            run = 0
            continue
        m = RE_DIALOG.match(line)
        if m:
            scene['lines'].append({
                'type': 'dialog',
                'speaker': m.group(1).strip(),
                'emotion': m.group(2) or '',
                'voice': m.group(3) or '',
                'text': m.group(4),
            })
            run = 0
            continue
        scene['lines'].append({'type': 'other', 'text': line})
        result['warnings'].append(f'{scene["scene_id"]} 未识别行: {line}')
        unrecognized(index, line)

    table_text = [cell.text.strip() for table in tables for row in table.rows
                  for cell in row.cells if cell.text.strip()]
    result['parse_report'] = {'nonempty_paragraphs': len(paras),
                              'unrecognized_count': len(unknown),
                              'max_consecutive_unrecognized': maximum_run,
                              'unrecognized': unknown,
                              'table_text_count': len(table_text)}
    source = path.resolve()
    result['source'] = {'path': str(source), 'sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
                        'format': source.suffix.lower().lstrip('.') or 'unknown'}
    return result


def quality_errors(data):
    errors = []
    if not data.get('episodes') or not any(ep.get('scenes') for ep in data.get('episodes', [])):
        errors.append('解析结果为 0 集或 0 场，不能继续')
    report = data.get('parse_report', {})
    total = report.get('nonempty_paragraphs', 0)
    if total and report.get('unrecognized_count', 0) / total > 0.05:
        errors.append('未识别段落超过非空段落总数的 5%，须修正规范格式后重新解析')
    if report.get('max_consecutive_unrecognized', 0) >= 10:
        errors.append('连续至少 10 行无法解析，须修正规范格式后重新解析')
    if report.get('table_text_count', 0):
        errors.append('DOCX 表格中存在正文，当前解析器不能保证覆盖；请转换为标准段落后重试')
    return errors


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    output = Path(sys.argv[2]).resolve()
    if output == Path(sys.argv[1]).resolve() or output.suffix.lower() != '.json':
        sys.exit('错误: 解析输出须为独立的 .json 文件，不能覆盖原始 DOCX')
    data = parse(sys.argv[1])
    errors = quality_errors(data)
    if errors:
        for error in errors:
            print('错误:', error)
        print(json.dumps(data['parse_report'], ensure_ascii=False, indent=2))
        sys.exit(1)
    with open(output, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=1)

    n_scenes = sum(len(e['scenes']) for e in data['episodes'])
    print(f'剧名: {data["title"]}')
    print(f'集数: {len(data["episodes"])}  场次: {n_scenes}')
    print(f'出场人物(人物行汇总): {"、".join(sorted({c for e in data["episodes"] for s in e["scenes"] for c in s["characters"]}))}')
    print(f'场景名(去重前): {len({s["location"] for e in data["episodes"] for s in e["scenes"]})} 个')
    for w in data['warnings']:
        print('警告:', w)


if __name__ == '__main__':
    main()
