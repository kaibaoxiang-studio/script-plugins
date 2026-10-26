"""模板及导出验收回归：检查行为，不依赖某一剧本的字句。"""
import copy
import tempfile
import zipfile
from pathlib import Path
import unittest
from openpyxl import Workbook
from build_xlsx import write_plan_sheet
from check_workbook import audit_workbook
from export_contract import build_plan, local_notes, sheet_plan


def fixture():
    scene_ids = ['2-1', '1-1']
    data = {'episodes': [{'episode': 1, 'title': '回忆', 'scenes': [
        {'scene_id': sid, 'location': '房间', 'int_ext': '内', 'time': '日', 'characters': ['甲'], 'lines': []}
        for sid in scene_ids]}]}
    states = [{'id': 'S01', 'name': '白衣', 'description': '白衣', 'scenes': scene_ids, 'evidence': '依据', 'note': ''},
              {'id': 'S02', 'name': '蓝衣', 'description': '蓝衣', 'scenes': ['2-1'], 'evidence': '依据', 'note': ''}]
    def record(line, state, timeline, sid='2-1'):
        return {'id': 'A'+str(line), 'scene_id': sid, 'start_line': line, 'end_line': line, 'state_id': state,
                'timeline': timeline, 'instance': '本体', 'segment': states[int(state[-1])-1]['name'],
                'basis': 'explicit', 'note': '待确认'}
    char = {'id': 'CHR-001', 'name': '甲', 'aliases': [], 'role': '主角', 'description': '固定外观',
            'states': states, 'appearances': [record(1, 'S01', '现实'), record(3, 'S02', '闪回'),
                      record(5, 'S01', '现实'), record(1, 'S01', '之后', '1-1')]}
    second = copy.deepcopy(char)
    second.update(id='CHR-002', name='乙', states=[copy.deepcopy(states[0])], appearances=[record(1, 'S01', '现实')])
    assets = {'characters': [second, char], 'props': [], 'scenes': [{'id': 'SCN-001', 'name': '房间',
       'int_ext': '内', 'times': ['日'], 'description': '房间', 'subscenes': [], 'states': [
           {'id': 'S01', 'name': '原状', 'description': '房间', 'scenes': scene_ids, 'evidence': '依据', 'note': ''},
           {'id': 'S02', 'name': '搬离', 'description': '家具仍在', 'scenes': ['2-1'], 'evidence': '依据', 'note': '未搬走家具'}]}]}
    return data, assets


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.data, self.assets = fixture()
        self.before = copy.deepcopy(self.assets)
        self.plan = build_plan(self.data, self.assets)
        self.workbook = Workbook(); self.workbook.remove(self.workbook.active)
        for sheet in self.plan['sheets']: write_plan_sheet(self.workbook, sheet)

    def test_plan_is_non_mutating_and_presentation_order_is_preserved(self):
        self.assertEqual(self.assets, self.before)
        rows = self.plan['sheets'][0]['rows']
        self.assertEqual(rows[0][1], '甲')
        stages = rows[0][11]
        self.assertLess(stages.index('2-1 L1-1'), stages.index('2-1 L5-5'))
        self.assertLess(stages.index('2-1 L5-5'), stages.index('1-1 L1-1'))
        breakdown = self.plan['sheets'][-1]['rows'][0][7]
        self.assertLess(breakdown.index('L1-1'), breakdown.index('L3-3'))
        self.assertLess(breakdown.index('L3-3'), breakdown.index('L5-5'))

    def test_repeated_notes_do_not_drag_nonadjacent_records_forward(self):
        records = [{'scene_id': '1-1', 'start_line': i, 'end_line': i, 'note': note}
                   for i, note in enumerate(['甲', '乙', '甲'], 1)]
        text = local_notes(records)
        self.assertLess(text.index('L1-1'), text.index('L2-2'))
        self.assertLess(text.index('L2-2'), text.index('L3-3'))

    def test_scene_state_note_and_breakdown_state_are_not_omitted(self):
        self.assertIn('未搬走家具', self.plan['sheets'][1]['rows'][0][7])
        self.assertIn('S02 搬离', self.plan['sheets'][-1]['rows'][0][6])
        self.assertIn('未搬走家具', self.plan['sheets'][-1]['rows'][0][12])

    def test_written_workbook_matches_contract(self):
        self.assertEqual(audit_workbook(self.workbook, self.plan), [])

    def test_content_loss_filter_and_group_style_are_detected(self):
        self.workbook['人物']['L2'] = '丢失阶段'
        self.workbook['道具'].auto_filter.ref = None
        self.workbook['场景']['A2'].fill = copy.copy(self.workbook['场景']['A1'].fill)
        errors = audit_workbook(self.workbook, self.plan)
        self.assertTrue(any('内容遗漏' in e for e in errors))
        self.assertTrue(any('筛选' in e for e in errors))
        self.assertTrue(any('底色' in e for e in errors))

    def test_same_line_state_order_respects_character_offsets(self):
        c = self.assets['characters'][1]
        c['appearances'][0].update(start_line=1, start_offset=20)
        c['appearances'][1].update(start_line=1, start_offset=2)
        plan = build_plan(self.data, self.assets)
        self.assertEqual(plan['sheets'][0]['rows'][0][5], 'CHR-001-S02')

    def test_overflow_is_blocked_instead_of_truncated(self):
        with self.assertRaises(ValueError):
            sheet_plan('人物', [['长文本' * 5000] + [''] * 12], [])

    def test_local_generic_register_name_cannot_be_global_alias(self):
        from prop_identity import validate_props
        from test_prop_identity import fixture as prop_fixture
        data, assets = prop_fixture()
        assets['props'][0]['aliases'].append('登记本')
        assets['props'][0]['alias_evidence'].append({'alias': '登记本', 'basis': 'continuity',
            'reason': '测试泛称', 'evidence_refs': [assets['props'][0]['mentions'][0]['source'], assets['props'][0]['mentions'][2]['source']]})
        errors, _ = validate_props(data, assets, check_derived=False)
        self.assertTrue(any('泛称「登记本」' in error for error in errors))

    def test_filter_patch_is_idempotent_and_preserves_other_zip_parts(self):
        from add_native_filters import add_filters
        with tempfile.TemporaryDirectory(prefix='asset-filter-test-') as directory:
            filename = Path(directory) / 'book.xlsx'
            self.workbook['人物'].auto_filter.ref = None
            self.workbook.save(filename)
            with zipfile.ZipFile(filename) as z:
                before = {item: z.read(item) for item in z.namelist() if not item.startswith('xl/worksheets/sheet')}
            add_filters(filename, self.plan)
            add_filters(filename, self.plan)
            from openpyxl import load_workbook
            self.assertEqual(audit_workbook(load_workbook(filename), self.plan), [])
            with zipfile.ZipFile(filename) as z:
                self.assertEqual(before, {item: z.read(item) for item in before})


if __name__ == '__main__': unittest.main()
