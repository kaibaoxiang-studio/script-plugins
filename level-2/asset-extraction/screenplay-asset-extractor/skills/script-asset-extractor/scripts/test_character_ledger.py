#!/usr/bin/env python3
"""人物状态归属回归验证：python3 scripts/test_character_ledger.py。"""
import copy
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
import docx

from build_xlsx import validate
from character_ledger import scene_index, sync_characters, validate_characters
from character_changes import change_candidates, covers
from warning_approvals import apply_approvals, warning_report
from show_character import candidates
from parse_script import parse


def fixture():
    def scene(sid, text, names=None):
        return {'scene_id': sid, 'location': '大厅', 'int_ext': '内', 'time': '日',
                'characters': ['莱拉'] if names is None else names,
                'lines': [{'type': 'action', 'text': t} for t in text]}

    scenes = [scene('1-1', ['莱拉穿白裙入场。', '莱拉的白裙染血。', '莱拉化为银狐。']),
              scene('1-2', ['莱拉恢复人形，白裙仍然染血。']),
              scene('1-3', ['闪回，童年莱拉穿蓝衣。']),
              scene('1-4', ['返回现实，莱拉接着沿走廊前进。']),
              scene('1-5', ['众人只谈起莱拉，并未呈现她的画面。'], [])]
    data = {'title': '回归剧本', 'episodes': [{'episode': 1, 'title': '测试', 'summary': '', 'scenes': scenes}]}

    def ref(sid, line):
        return {'scene_id': sid, 'line': line,
                'quote': next(s for s in scenes if s['scene_id'] == sid)['lines'][line - 1]['text']}

    def state(sid, name, source, costume, form='未说明'):
        return {'id': sid, 'name': name, 'basis': 'explicit',
                'description': f'{name}，其余外观未说明。',
                'visual': {'age': '未说明', 'costume': costume, 'form': form,
                           'hair_makeup': '未说明', 'injury': '未说明', 'stains': '未说明'},
                'evidence_refs': [source], 'scenes': [], 'evidence': '', 'note': ''}

    states = [state('S01', '白裙', ref('1-1', 1), '白裙'),
              state('S02', '白裙染血', ref('1-1', 2), '白裙染血'),
              state('S03', '银狐', ref('1-1', 3), '未说明', '银狐'),
              state('S04', '童年蓝衣', ref('1-3', 1), '蓝衣')]
    states[1]['evidence_refs'].append(ref('1-2', 1))

    def record(rid, sid, line, stid, timeline='现实', basis='explicit', boundary='continuous', source=None):
        return {'id': rid, 'scene_id': sid, 'start_line': line, 'end_line': line,
                'segment': rid, 'timeline': timeline, 'instance': '本体', 'state_id': stid,
                'basis': basis, 'boundary': boundary, 'from': source,
                'continuity': '现实线的走廊行动在插入闪回后继续，无现实时间跳跃。' if basis == 'inherited' else '',
                'presence': {'kind': 'header', 'refs': []},
                'evidence_refs': [ref(sid, line)] if basis == 'explicit' else [], 'note': ''}

    appearances = [record('A001', '1-1', 1, 'S01', boundary='reset'),
                   record('A002', '1-1', 2, 'S02'), record('A003', '1-1', 3, 'S03'),
                   record('A004', '1-2', 1, 'S02'),
                   record('A005', '1-3', 1, 'S04', timeline='童年闪回', boundary='reset'),
                   record('A006', '1-4', 1, 'S02', basis='inherited', source='A004')]
    character = {'id': 'CHR-001', 'name': '莱拉', 'aliases': [], 'role': '主角',
                 'description': '未说明的固定外观待确认。', 'states': states, 'appearances': appearances}
    assets = {'characters': [character],
              'scenes': [{'id': 'SCN-001', 'name': '大厅', 'int_ext': '内', 'times': ['日'],
                          'states': [{'id': 'S01', 'name': '常态', 'scenes': [s['scene_id'] for s in scenes]}]}],
              'props': []}
    record_fixture_reviews(data, assets)
    return data, assets


def record_fixture_reviews(data, assets):
    """仅给确定的测试画面写复核记录；产品同步函数不生成判读结论。"""
    for char in assets['characters']:
        reviews = []
        for candidate in change_candidates(data, char):
            records = [rec for rec in char['appearances'] if covers(candidate, rec, scene_index(data))]
            if not records:
                raise AssertionError('测试原文候选缺少对应阶段')
            reviews.append({'candidate_id': candidate['candidate_id'],
                            'source': candidate['source'], 'decision': 'recorded',
                            'appearance_ids': [rec['id'] for rec in records],
                            'reason': '测试画面已在关联阶段明确呈现该外观。'})
        char['change_reviews'] = reviews


def add_warning_prop(data, assets):
    text = '一名路人拿着匕首站在门口。'
    data['episodes'][0]['scenes'][-1]['lines'] = [{'type': 'action', 'text': text}]
    low = text.index('匕首')
    assets['props'] = [{'id': 'PRP-001', 'name': '匕首', 'owner': '莱拉', 'aliases': [],
                        'alias_evidence': [], 'states': [{'id': 'S01', 'name': '常态', 'scenes': ['1-5']}],
                        'mentions': [{'id': 'M001', 'raw_name': '匕首', 'basis': 'named',
                                      'visibility': 'visible', 'timeline': '现实', 'state_id': 'S01',
                                      'start_offset': low, 'end_offset': low + 2,
                                      'source': {'scene_id': '1-5', 'line': 1, 'quote': text},
                                      'holder': '路人', 'note': ''}]}]


class LedgerTests(unittest.TestCase):
    def setUp(self):
        self.data, self.assets = fixture()
        self.char = self.assets['characters'][0]

    def sync(self):
        errors, _ = sync_characters(self.data, self.assets)
        self.assertEqual(errors, [])

    def rejected(self, phrase):
        errors, _ = validate_characters(self.data, self.assets, check_derived=False)
        self.assertTrue(any(phrase in error for error in errors), errors)

    def test_same_scene_changes_and_flashback_return(self):
        self.sync()
        self.assertEqual(validate(self.data, self.assets), ([], []))
        self.assertEqual(self.char['states'][1]['scenes'], ['1-1', '1-2', '1-4'])
        self.assertEqual(self.char['states'][3]['scenes'], ['1-3'])
        self.assertNotIn('1-5', self.char['states'][1]['scenes'])

    def test_sync_is_idempotent(self):
        self.sync()
        original = copy.deepcopy(self.assets)
        self.sync()
        self.assertEqual(self.assets, original)

    def test_same_version_returns_in_one_scene_without_duplicate_scene(self):
        scene = self.data['episodes'][0]['scenes'][0]
        scene['lines'].append({'type': 'action', 'text': '莱拉换回干净的白裙。'})
        ref = {'scene_id': '1-1', 'line': 4, 'quote': '莱拉换回干净的白裙。'}
        self.char['states'][0]['evidence_refs'].append(copy.deepcopy(ref))
        rec = copy.deepcopy(self.char['appearances'][0])
        rec.update(id='A007', start_line=4, end_line=4, boundary='continuous', evidence_refs=[ref])
        self.char['appearances'].insert(3, rec)
        record_fixture_reviews(self.data, self.assets)
        self.sync()
        self.assertEqual(validate(self.data, self.assets), ([], []))
        self.assertEqual(self.char['states'][0]['scenes'], ['1-1'])
        self.assertEqual(sum(r['state_id'] == 'S01' for r in self.char['appearances']), 2)

    def test_docx_parse_preserves_line_positions_and_voice_candidates(self):
        with tempfile.TemporaryDirectory(prefix='character-parser-') as temp:
            path = Path(temp) / 'script.docx'
            document = docx.Document()
            for text in ['行号测试', '第1集｜开端', '1-1 大厅 内 日', '人物：莱拉',
                         '▲莱拉穿白裙入场。', '【闪回】', '莱拉[VO]：我记得那天。']:
                document.add_paragraph(text)
            document.save(path)
            parsed = parse(path)
            lines = parsed['episodes'][0]['scenes'][0]['lines']
            self.assertEqual(lines[0]['text'], '莱拉穿白裙入场。')
            self.assertEqual(lines[1], {'type': 'other', 'text': '【闪回】'})
            self.assertEqual(lines[2]['voice'], 'VO')
            hits = list(candidates(parsed, ['莱拉']))
            self.assertIn('对白/画外音（须判读）', hits[0][2])

    def test_manual_scene_corruption_rejected(self):
        self.sync()
        self.char['states'][0]['scenes'].append('1-3')
        errors, _ = validate(self.data, self.assets)
        self.assertTrue(any('scenes 与逐场记录不一致' in error for error in errors))

    def test_false_quote_rejected(self):
        self.char['appearances'][0]['evidence_refs'][0]['quote'] = '莱拉身穿红色婚纱'
        self.rejected('引文不在指定原文中')

    def test_wrong_scene_evidence_rejected(self):
        self.char['appearances'][0]['evidence_refs'] = self.char['states'][3]['evidence_refs']
        self.rejected('外观依据必须位于本场本阶段')

    def test_phase_overlap_rejected(self):
        self.char['appearances'][1]['start_line'] = 1
        self.rejected('阶段区间重叠')

    def test_wrong_timeline_inheritance_rejected(self):
        self.char['appearances'][-1]['from'] = 'A005'
        self.rejected('继承必须指向')

    def test_skipping_recent_appearance_rejected(self):
        self.char['appearances'][-1]['from'] = 'A002'
        self.rejected('继承必须指向')

    def test_reset_cannot_inherit(self):
        self.char['appearances'][-1]['boundary'] = 'reset'
        self.rejected('继承必须指向')

    def test_inheritance_requires_continuity_reason(self):
        self.char['appearances'][-1]['continuity'] = ''
        self.rejected('连续性判读说明')

    def test_legacy_data_requires_source_review(self):
        del self.char['appearances']
        self.rejected('旧版数据须回查原文迁移')

    def test_duplicate_state_rejected(self):
        self.char['states'].append(copy.deepcopy(self.char['states'][0]))
        self.rejected('状态ID为空或重复')

    def test_duplicate_scene_rejected(self):
        self.data['episodes'][0]['scenes'].append(copy.deepcopy(self.data['episodes'][0]['scenes'][0]))
        self.rejected('剧本场次号重复')

    def test_action_presence_valid_when_header_omits_actor(self):
        self.data['episodes'][0]['scenes'][1]['characters'] = []
        rec = self.char['appearances'][3]
        rec['presence'] = {'kind': 'action', 'refs': copy.deepcopy(rec['evidence_refs'])}
        self.sync()
        self.assertEqual(validate(self.data, self.assets), ([], []))

    def test_dialogue_not_accepted_as_action_presence(self):
        scene = self.data['episodes'][0]['scenes'][1]
        scene['characters'] = []
        scene['lines'][0]['type'] = 'dialog'
        scene['lines'][0]['speaker'] = '旁白'
        rec = self.char['appearances'][3]
        rec['presence'] = {'kind': 'action', 'refs': copy.deepcopy(rec['evidence_refs'])}
        self.rejected('本阶段的动作行')

    def test_candidate_mentions_are_not_auto_assignments(self):
        hits = list(candidates(self.data, ['莱拉']))
        self.assertEqual(len(hits), 5)
        self.assertEqual(hits[-1][2], ['正文提及（须判读）'])
        self.sync()
        self.assertFalse(any('1-5' in s['scenes'] for s in self.char['states']))

    def test_unknown_appearance_is_valid_and_not_invented_costume(self):
        self.char['states'].append({'id': 'S05', 'name': '外观未说明', 'basis': 'unspecified',
                                   'description': '当前外观未说明。', 'evidence_refs': [],
                                   'visual': {key: '未说明' for key in self.char['states'][0]['visual']},
                                   'scenes': [], 'evidence': '', 'note': '时间跳跃后的服装未说明。'})
        rec = self.char['appearances'][-1]
        rec.update(state_id='S05', basis='unspecified', boundary='reset', **{'from': None})
        self.sync()
        self.assertEqual(validate(self.data, self.assets), ([], []))

    def test_single_line_changes_use_offsets(self):
        scene = self.data['episodes'][0]['scenes'][0]
        texts = [line['text'] for line in scene['lines']]
        scene['lines'] = [{'type': 'action', 'text': ''.join(texts)}]
        cursor = 0
        for i, text in enumerate(texts):
            rec = self.char['appearances'][i]
            rec.update(start_line=1, end_line=1, start_offset=cursor, end_offset=cursor + len(text))
            rec['evidence_refs'][0]['line'] = 1
            self.char['states'][i]['evidence_refs'][0]['line'] = 1
            cursor += len(text)
        record_fixture_reviews(self.data, self.assets)
        self.sync()
        self.assertEqual(validate(self.data, self.assets), ([], []))

    def test_quote_outside_offsets_rejected(self):
        rec = self.char['appearances'][0]
        rec['start_offset'] = 2
        self.rejected('引文超出本阶段字符区间')

    def test_omitted_change_cannot_pass(self):
        self.char['appearances'].pop(1)
        self.char['appearances'][0]['end_line'] = 2
        self.char['states'][1]['evidence_refs'].pop(0)
        self.char['change_reviews'] = [review for review in self.char['change_reviews']
                                     if review['source']['scene_id'] != '1-1' or review['source']['line'] != 2]
        errors, _ = sync_characters(self.data, self.assets)
        self.assertTrue(any('未复核外观变化候选' in error and '染血' in error for error in errors), errors)

    def test_stale_change_review_does_not_pass_after_text_edit(self):
        self.data['episodes'][0]['scenes'][0]['lines'][0]['text'] += '她停下来。'
        self.rejected('变化候选不存在或原文已修改')

    def test_recorded_change_must_have_local_visual_evidence(self):
        review = next(r for r in self.char['change_reviews']
                      if r['source']['scene_id'] == '1-1' and r['source']['line'] == 2)
        review['appearance_ids'] = ['A001']
        self.char['appearances'].pop(1)
        self.char['states'][1]['evidence_refs'].pop(0)
        self.char['appearances'][0]['end_line'] = 2
        self.rejected('已记录变化缺少该阶段的明确外观依据')

    def test_false_positive_can_be_reviewed_without_inventing_appearance(self):
        self.data['episodes'][0]['scenes'][-1]['lines'][0]['text'] = '众人谈起莱拉计划换上婚纱，但她没有出现在画面中。'
        candidate = list(change_candidates(self.data, self.char))[-1]
        self.char['change_reviews'].append({'candidate_id': candidate['candidate_id'],
                                          'source': candidate['source'], 'decision': 'not_visible',
                                          'appearance_ids': [], 'reason': '仅描述计划，未出镜。'})
        self.sync()
        self.assertEqual(validate(self.data, self.assets), ([], []))
        self.assertFalse(any(r['scene_id'] == '1-5' for r in self.char['appearances']))

    def test_invalid_field_types_return_errors_instead_of_crashing(self):
        for key, value in [('evidence_refs', 7), ('scenes', None), ('visual', []), ('id', {})]:
            with self.subTest(key=key):
                data, assets = fixture()
                assets['characters'][0]['states'][0][key] = value
                errors, _ = validate(data, assets)
                self.assertTrue(any('类型错误' in error for error in errors), errors)
                errors, _ = sync_characters(data, assets)
                self.assertTrue(errors)

    def test_appearance_notes_exported_to_both_sheets(self):
        from openpyxl import Workbook
        from build_xlsx import write_asset_sheet, write_breakdown
        marker = '鞋子未说明，待设计确认'
        self.char['appearances'][-1]['note'] = marker
        self.sync()
        wb = Workbook()
        write_asset_sheet(wb, 'characters', '人物', '人物名', [('role', '角色类型')], self.assets)
        write_breakdown(wb, self.data, self.assets)
        self.assertIn(marker, wb['人物'].cell(3, 11).value)
        self.assertIn('1-4', wb['人物'].cell(3, 11).value)
        self.assertIn(marker, wb['场次总表'].cell(5, 13).value)

    def test_exact_warning_approval_expires_with_source_change(self):
        self.sync()
        warnings = ['1-5 道具归属人物不在场']
        report = warning_report(self.data, self.assets, warnings)
        report['approvals'] = [dict(report['warnings'][0], reason='他人持有，已回查原文。', confirmed_by='用户')]
        self.assertEqual(apply_approvals(self.data, self.assets, warnings, report)[0], [])
        self.char['appearances'][-1]['note'] = '新修改'
        self.assertTrue(apply_approvals(self.data, self.assets, warnings, report)[2])

    def test_approval_requires_reason_and_cannot_cover_different_warning(self):
        self.sync()
        warnings = ['1-5 道具归属人物不在场', '1-2 另一条警告']
        report = warning_report(self.data, self.assets, warnings)
        report['approvals'] = [dict(report['warnings'][0], reason='他人持有。', confirmed_by='用户')]
        self.assertEqual(apply_approvals(self.data, self.assets, warnings, report)[0], ['1-2 另一条警告'])
        report['approvals'][0]['reason'] = ''
        self.assertTrue(apply_approvals(self.data, self.assets, warnings, report)[2])

    def test_character_frequency_counts_distinct_scenes_and_keeps_groups(self):
        from openpyxl import Workbook
        from build_xlsx import write_asset_sheet
        self.sync()
        rare = copy.deepcopy(self.char)
        rare.update(id='CHR-002', name='测试人物', states=[copy.deepcopy(self.char['states'][0])])
        rare['appearances'] = [copy.deepcopy(self.char['appearances'][0]) for _ in range(20)]
        assets = copy.deepcopy(self.assets)
        assets['characters'] = [rare, self.char]
        before = copy.deepcopy(assets)
        wb = Workbook()
        write_asset_sheet(wb, 'characters', '人物', '人物名', [('role', '角色类型')], assets, self.data)
        sheet = wb['人物']
        names = [row[1].value for row in sheet.iter_rows(min_row=2) if row[1].value]
        self.assertEqual(names, ['莱拉', '测试人物'])
        self.assertEqual(sheet.cell(2, 1).value, 'CHR-001')
        self.assertEqual(sheet.cell(6, 1).value, 'CHR-002')
        self.assertIn('A2:A5', [str(area) for area in sheet.merged_cells.ranges])
        self.assertEqual(assets, before)

    def test_state_rows_follow_first_use_including_same_scene_phases(self):
        from openpyxl import Workbook
        from build_xlsx import write_asset_sheet
        self.sync()
        self.char['states'].reverse()
        before = copy.deepcopy(self.assets)
        wb = Workbook()
        write_asset_sheet(wb, 'characters', '人物', '人物名', [('role', '角色类型')], self.assets, self.data)
        ids = [row[5].value for row in wb['人物'].iter_rows(min_row=2)]
        self.assertEqual(ids, ['CHR-001-S01', 'CHR-001-S02', 'CHR-001-S03', 'CHR-001-S04'])
        self.assertEqual(self.assets, before)

    def test_state_order_uses_script_presentation_not_numeric_scene_id(self):
        from openpyxl import Workbook
        from build_xlsx import write_asset_sheet
        scenes = self.data['episodes'][0]['scenes']
        scenes.insert(0, scenes.pop(2))
        order = {scene['scene_id']: i for i, scene in enumerate(scenes)}
        self.char['appearances'].sort(key=lambda rec: (order[rec['scene_id']], rec['start_line']))
        self.sync()
        self.assertEqual(validate(self.data, self.assets), ([], []))
        wb = Workbook()
        write_asset_sheet(wb, 'characters', '人物', '人物名', [('role', '角色类型')], self.assets, self.data)
        ids = [row[5].value for row in wb['人物'].iter_rows(min_row=2)]
        self.assertEqual(ids, ['CHR-001-S04', 'CHR-001-S01', 'CHR-001-S02', 'CHR-001-S03'])

    def test_cli_invalid_inputs_do_not_traceback_or_overwrite(self):
        self.sync()
        with tempfile.TemporaryDirectory(prefix='character-bad-input-') as temp:
            root = Path(temp)
            source, assets, output = root / 'scenes.json', root / 'assets.json', root / 'output.xlsx'
            source.write_text(json.dumps(self.data, ensure_ascii=False), encoding='utf-8')
            self.char['states'][0]['evidence_refs'] = 7
            original = json.dumps(self.assets, ensure_ascii=False)
            assets.write_text(original, encoding='utf-8')
            command = [sys.executable, str(Path(__file__).with_name('build_xlsx.py')), str(source), str(assets)]
            for mode in ['--check', '--sync-character-scenes', str(output)]:
                result = subprocess.run(command + [mode], capture_output=True, text=True)
                self.assertNotEqual(result.returncode, 0)
                self.assertIn('类型错误', result.stdout)
                self.assertNotIn('Traceback', result.stdout + result.stderr)
                self.assertFalse(output.exists())
                self.assertEqual(assets.read_text(encoding='utf-8'), original)

    def test_cli_approval_cannot_override_error(self):
        self.sync()
        with tempfile.TemporaryDirectory(prefix='character-approval-errors-') as temp:
            root = Path(temp)
            add_warning_prop(self.data, self.assets)
            self.char['states'][0]['scenes'].append('1-5')
            errors, warnings = validate(self.data, self.assets)
            self.assertTrue(errors)
            report = warning_report(self.data, self.assets, warnings)
            report['approvals'] = [dict(item, reason='测试确认该条警告。', confirmed_by='用户') for item in report['warnings']]
            source, assets, approvals, output = [root / name for name in ['scenes.json', 'assets.json', 'approvals.json', 'blocked.xlsx']]
            for path, value in [(source, self.data), (assets, self.assets), (approvals, report)]:
                path.write_text(json.dumps(value, ensure_ascii=False), encoding='utf-8')
            command = [sys.executable, str(Path(__file__).with_name('build_xlsx.py')), str(source), str(assets), str(output), '--warning-approvals', str(approvals)]
            result = subprocess.run(command, capture_output=True, text=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertFalse(output.exists())

    def test_cli_export_and_warning_gate(self):
        self.sync()
        with tempfile.TemporaryDirectory(prefix='character-ledger-') as temp:
            root = Path(temp)
            scene_path, asset_path, output = root / 'scenes.json', root / 'assets.json', root / 'assets.xlsx'
            scene_path.write_text(json.dumps(self.data, ensure_ascii=False), encoding='utf-8')
            asset_path.write_text(json.dumps(self.assets, ensure_ascii=False), encoding='utf-8')
            command = [sys.executable, str(Path(__file__).with_name('build_xlsx.py')), str(scene_path), str(asset_path)]
            for mode in ['--sync-character-scenes', '--check', str(output)]:
                result = subprocess.run(command + [mode], capture_output=True, text=True)
                self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            workbook = load_workbook(output)
            self.assertEqual(workbook.sheetnames, ['人物', '场景', '道具', '场次总表'])
            self.assertEqual(workbook['人物'].cell(1, 12).value, '场次·阶段·时间线')
            breakdown = workbook['场次总表'].cell(2, 8).value
            self.assertLess(breakdown.index('白裙 ['), breakdown.index('白裙染血 ['))
            self.assertLess(breakdown.index('白裙染血 ['), breakdown.index('银狐 ['))
            self.assertIn('童年闪回', workbook['场次总表'].cell(4, 8).value)
            workbook.close()

            add_warning_prop(self.data, self.assets)
            scene_path.write_text(json.dumps(self.data, ensure_ascii=False), encoding='utf-8')
            asset_path.write_text(json.dumps(self.assets, ensure_ascii=False), encoding='utf-8')
            blocked = root / 'blocked.xlsx'
            result = subprocess.run(command + [str(blocked)], capture_output=True, text=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn('警告', result.stdout)
            self.assertFalse(blocked.exists())
            result = subprocess.run(command + ['--warning-report'], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            report = json.loads(result.stdout)
            report['approvals'] = [dict(report['warnings'][0], reason='原文由他人临时持有，测试用户确认。', confirmed_by='用户')]
            approval_path = root / 'approvals.json'
            approval_path.write_text(json.dumps(report, ensure_ascii=False), encoding='utf-8')
            result = subprocess.run(command + [str(blocked), '--warning-approvals', str(approval_path)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertIn('未确认警告 0 条，已确认警告 1 条', result.stdout)
            workbook = load_workbook(blocked)
            self.assertIn('测试用户确认', workbook['场次总表'].cell(6, 11).value)
            workbook.close()


if __name__ == '__main__':
    unittest.main(verbosity=2)
