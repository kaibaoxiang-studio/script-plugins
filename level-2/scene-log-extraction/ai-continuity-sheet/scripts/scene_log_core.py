"""Scene-log data contract, source tracing, and Excel export.

The agent chooses beat boundaries and proposes directing notes. This module owns
lossless source mapping, structural checks, and the published artifacts.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HEADERS = [
    "集数", "场次", "片段", "出场人物", "场景", "道具", "摄影", "灯光/色卡",
    "台词", "时长", "微表情情绪指导", "动作表演指导", "剧情大纲",
]
SUGGESTION_KEYS = {
    "camera": "摄影", "lighting": "灯光/色卡", "micro_expression": "微表情情绪指导",
    "action_performance": "动作表演指导", "synopsis": "剧情大纲",
}
SCRIPT_CELL_LIMIT = 32000


def _unique(items: list[str]) -> list[str]:
    return list(dict.fromkeys(item for item in items if item))


def load_source(path: str | Path) -> dict[str, Any]:
    source = Path(path)
    if not source.is_file():
        raise ValueError(f"找不到一级标准化 JSON：{source}")
    try:
        data = json.loads(source.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise ValueError(f"无法读取一级标准化 JSON：{exc}") from exc
    if not isinstance(data, dict) or data.get("schema_version") != "standardized_screenplay_v1":
        raise ValueError("输入必须是一级产出的 standardized_screenplay_v1 JSON")
    episodes = data.get("episodes")
    if not isinstance(episodes, list) or not episodes:
        raise ValueError("一级 JSON 没有分集；请先完成剧本标准化或细纲创作")
    scene_count = 0
    seen_episode_numbers = set()
    seen_scene_ids = set()
    for ep_index, episode in enumerate(episodes, 1):
        if not isinstance(episode, dict) or not isinstance(episode.get("scenes"), list):
            raise ValueError(f"第 {ep_index} 集缺少 scenes 数组")
        episode_number = episode.get("number", ep_index)
        if not isinstance(episode_number, int) or episode_number in seen_episode_numbers:
            raise ValueError(f"第 {ep_index} 集集号无效或重复")
        seen_episode_numbers.add(episode_number)
        for scene_index, scene in enumerate(episode["scenes"], 1):
            if not isinstance(scene, dict) or not isinstance(scene.get("blocks"), list) or not scene["blocks"]:
                raise ValueError(f"第 {ep_index} 集第 {scene_index} 场没有正文 blocks；请先补全一级剧本")
            scene_id = scene.get("id", f"scn-{scene_count + 1:03d}")
            if not isinstance(scene_id, str) or scene_id in seen_scene_ids:
                raise ValueError(f"第 {ep_index} 集第 {scene_index} 场 scene id 无效或重复")
            seen_scene_ids.add(scene_id)
            for block_index, block in enumerate(scene["blocks"], 1):
                if not isinstance(block, dict) or not isinstance(block.get("text"), str) or not block["text"]:
                    raise ValueError(f"第 {ep_index} 集第 {scene_index} 场第 {block_index} 块正文缺失")
                if not isinstance(block.get("source_refs"), list) or not block["source_refs"]:
                    raise ValueError(f"第 {ep_index} 集第 {scene_index} 场第 {block_index} 块缺少来源位置")
            scene_count += 1
    if not scene_count:
        raise ValueError("一级 JSON 没有可处理的场景；若为纯细纲，请先完成一级剧本创作")
    return data


def _action_fragments(text: str) -> list[str]:
    # Split only at full sentence boundaries; concatenation is byte-for-byte
    # identical to the source block. The agent can regroup adjacent units.
    parts = []
    start = depth = 0
    for index, char in enumerate(text):
        if char in "“‘":
            depth += 1
        elif char in "”’":
            depth = max(0, depth - 1)
        boundary = (char in "。！？!?" and depth == 0) or (
            char in "”’" and depth == 0 and index > 0 and text[index - 1] in "。！？!?"
        )
        if boundary:
            parts.append(text[start:index + 1])
            start = index + 1
    if start < len(text):
        parts.append(text[start:])
    return parts or [text]


def prepare_scenes(data: dict[str, Any], selected: list[int] | None = None) -> list[dict[str, Any]]:
    known = {int(ep.get("number", index)) for index, ep in enumerate(data["episodes"], 1)}
    if selected and not set(selected) <= known:
        raise ValueError(f"指定集不存在：{sorted(set(selected) - known)}")
    scenes: list[dict[str, Any]] = []
    absolute_scene = 0
    for ep_index, episode in enumerate(data["episodes"], 1):
        episode_number = int(episode.get("number", ep_index))
        for scene_index, scene in enumerate(episode["scenes"], 1):
            absolute_scene += 1
            if selected and episode_number not in selected:
                continue
            units = []
            for block_index, block in enumerate(scene["blocks"]):
                parts = _action_fragments(block["text"]) if block.get("type") == "action" else [block["text"]]
                assert "".join(parts) == block["text"]
                for fragment_index, text in enumerate(parts):
                    units.append({
                        "id": f"u{len(units) + 1:03d}", "block_index": block_index,
                        "fragment_index": fragment_index, "type": block.get("type", "action"),
                        "text": text, "source_refs": block["source_refs"],
                        "origin": block.get("origin", "source"),
                        "speaker": block.get("speaker"), "voice": block.get("voice"),
                        "language": block.get("language"),
                    })
            scenes.append({
                "episode_id": episode.get("id", f"ep-{ep_index:03d}"),
                "episode_number": episode_number,
                "episode_title": episode.get("title", ""),
                "scene_id": scene.get("id", f"scn-{absolute_scene:03d}"),
                "scene_number": scene.get("number", f"{episode_number}-{scene_index}"),
                "sc_number": f"SC{absolute_scene:02d}",
                "location": scene.get("location", ""),
                "interior_exterior": scene.get("interior_exterior", ""),
                "time": scene.get("time", ""),
                "characters": scene.get("characters", []),
                "source_refs": scene.get("source_refs", []),
                "units": units,
            })
    if not scenes:
        raise ValueError("指定集没有可处理的场景正文；请先补全一级剧本")
    return scenes


def validate_groups(scene: dict[str, Any], groups: list[dict[str, Any]]) -> None:
    expected = [unit["id"] for unit in scene["units"]]
    actual = []
    for index, group in enumerate(groups, 1):
        ids = group.get("unit_ids") if isinstance(group, dict) else None
        if not isinstance(ids, list) or not ids or not all(isinstance(item, str) for item in ids):
            raise ValueError(f"{scene['scene_id']} 第 {index} 拍没有有效 unit_ids")
        actual.extend(ids)
    if actual != expected:
        raise ValueError(f"{scene['scene_id']} 分拍未按原顺序完整覆盖正文；预期 {expected}，实际 {actual}")


def dialogue_for_units(units: list[dict[str, Any]]) -> tuple[str, list[dict[str, Any]]]:
    items = []
    displays = []
    for unit in units:
        if unit["type"] != "dialogue":
            continue
        item = {key: unit.get(key) for key in ("text", "speaker", "voice", "language", "source_refs", "origin")}
        items.append(item)
        voice = f"（{unit['voice']}）" if unit.get("voice") else ""
        speaker = unit.get("speaker") or "说话人未标明"
        displays.append(f"{speaker}{voice}：{unit['text']}")
    return "\n".join(displays) if displays else "—", items


def _suggestion(text: str, rationale: str, refs: list[str]) -> dict[str, Any]:
    return {"text": text or "—", "review_status": "pending", "rationale": rationale or "待人工核对", "source_refs": refs}


def assemble(data: dict[str, Any], scenes: list[dict[str, Any]], results: dict[str, dict[str, Any]]) -> dict[str, Any]:
    episode_map: dict[int, dict[str, Any]] = {}
    for scene in scenes:
        result = results.get(scene["scene_id"])
        if not isinstance(result, dict):
            raise ValueError(f"缺少 {scene['scene_id']} 的分场结果")
        segmentation = result.get("segmentation", {})
        enriched = result.get("enrichment", {})
        groups = segmentation.get("groups", [])
        validate_groups(scene, groups)
        notes = enriched.get("beats", [])
        if len(notes) != len(groups):
            raise ValueError(f"{scene['scene_id']} 建议条数与分拍数不一致")
        units_by_id = {unit["id"]: unit for unit in scene["units"]}
        beats = []
        for index, (group, note) in enumerate(zip(groups, notes), 1):
            if note.get("unit_ids") != group["unit_ids"]:
                raise ValueError(f"{scene['scene_id']} P{index:02d} 建议未对应原分拍")
            units = [units_by_id[unit_id] for unit_id in group["unit_ids"]]
            refs = _unique([ref for unit in units for ref in unit["source_refs"]])
            dialogue, dialogue_blocks = dialogue_for_units(units)
            characters = note.get("characters", [])
            if not isinstance(characters, list):
                raise ValueError(f"{scene['scene_id']} P{index:02d} 出场人物格式错误")
            characters = _unique([str(item) for item in characters])
            allowed_characters = set(scene["characters"]) | {unit["speaker"] for unit in units if unit.get("speaker")}
            unexpected = set(characters) - allowed_characters
            if unexpected:
                raise ValueError(f"{scene['scene_id']} P{index:02d} 出场人物没有场景或说话人依据：{sorted(unexpected)}")
            props = note.get("props", [])
            if not isinstance(props, list):
                raise ValueError(f"{scene['scene_id']} P{index:02d} 道具格式错误")
            props = _unique([str(item) for item in props])
            rationale = note.get("rationale", "")
            suggestions = {key: _suggestion(str(note.get(key, "")), rationale, refs) for key in SUGGESTION_KEYS}
            duration = note.get("duration_estimate", {})
            seconds = duration.get("seconds") if isinstance(duration, dict) else None
            if not isinstance(seconds, (int, float)) or isinstance(seconds, bool) or seconds <= 0:
                raise ValueError(f"{scene['scene_id']} P{index:02d} 时长建议无有效秒数")
            duration_value = {"seconds": seconds, "basis": str(duration.get("basis", "")), "review_status": "pending", "source_refs": refs}
            fields = {
                "集数": f"EP{scene['episode_number']:02d}", "场次": scene["sc_number"],
                "片段": f"P{index:02d}", "出场人物": "、".join(characters) or "—",
                "场景": " ".join(filter(None, [scene["location"], scene["interior_exterior"], scene["time"]])) or "—",
                "道具": "、".join(props) or "—", "摄影": suggestions["camera"]["text"],
                "灯光/色卡": suggestions["lighting"]["text"], "台词": dialogue,
                "时长": f"约{seconds:g}s", "微表情情绪指导": suggestions["micro_expression"]["text"],
                "动作表演指导": suggestions["action_performance"]["text"],
                "剧情大纲": suggestions["synopsis"]["text"],
            }
            beats.append({
                "id": f"{scene['scene_id']}/P{index:02d}", "beat_id": f"B{index:02d}",
                "segment_id": f"P{index:02d}", "unit_ids": group["unit_ids"],
                "source_units": units, "source_refs": refs, "dialogue_blocks": dialogue_blocks,
                "source_origin": "generated" if any(unit["origin"] == "generated" for unit in units) else "source",
                "fields": fields, "suggestions": suggestions, "duration_estimate": duration_value,
                "uncertainties": note.get("uncertainties", []),
                "beat_reason": group.get("reason", ""),
            })
        ep = episode_map.setdefault(scene["episode_number"], {
            "id": scene["episode_id"], "number": scene["episode_number"],
            "title": scene["episode_title"], "scenes": [],
        })
        ep["scenes"].append({
            "id": scene["scene_id"], "number": scene["scene_number"],
            "sc_number": scene["sc_number"], "location": scene["location"],
            "interior_exterior": scene["interior_exterior"], "time": scene["time"],
            "source_refs": scene["source_refs"], "beats": beats,
        })
    return {"schema_version": "scene_log_v1", "source_schema_version": data["schema_version"],
            "source_manifest": data.get("source_manifest", {}),
            "source_coverage": data.get("coverage", {}),
            "episodes": [episode_map[number] for number in sorted(episode_map)]}


def episode_script_text(source_episode: dict[str, Any]) -> str:
    lines = [f"第 {source_episode.get('number')} 集｜{source_episode.get('title', '')}"]
    if source_episode.get("summary"):
        lines.append(f"本集概要：{source_episode['summary']}")
    for scene in source_episode.get("scenes", []):
        lines.append(scene.get("source_heading") or scene.get("number", ""))
        if scene.get("characters"):
            lines.append("人物：" + "、".join(scene["characters"]))
        for block in scene.get("blocks", []):
            if block.get("type") == "dialogue":
                voice = f"（{block['voice']}）" if block.get("voice") else ""
                lines.append(f"{block.get('speaker', '')}{voice}：{block['text']}")
            else:
                lines.append(block["text"])
    return "\n".join(lines)


def _chunks(text: str, limit: int = SCRIPT_CELL_LIMIT) -> list[str]:
    return [text[index:index + limit] for index in range(0, len(text), limit)] or [""]


def write_workbook(log: dict[str, Any], source: dict[str, Any], path: str | Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    by_number = {int(ep.get("number", index)): ep for index, ep in enumerate(source["episodes"], 1)}
    border = Border(bottom=Side(style="hair", color="D6DEE4"))
    for episode in log["episodes"]:
        ws = wb.create_sheet(f"EP{episode['number']:02d}")
        script = episode_script_text(by_number[episode["number"]])
        for chunk_index, chunk in enumerate(_chunks(script), 1):
            label = f"一级标准化剧本正文（{chunk_index}）\n" if len(script) > SCRIPT_CELL_LIMIT else "一级标准化剧本正文\n"
            ws.append([label + chunk] + [None] * 12)
            row = ws.max_row
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=13)
            cell = ws.cell(row, 1)
            cell.fill = PatternFill("solid", fgColor="FBF7E9")
            cell.font = Font(name="微软雅黑", size=10, color="4A3B2A")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[row].height = min(400, max(70, len(chunk) // 100 * 12))
        header_row = ws.max_row + 1
        ws.append(HEADERS)
        for col, header in enumerate(HEADERS, 1):
            cell = ws.cell(header_row, col)
            cell.fill = PatternFill("solid", fgColor="2C3E50")
            cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=10)
            cell.alignment = Alignment(wrap_text=True, horizontal="center")
            ws.column_dimensions[get_column_letter(col)].width = 30 if col > 3 else 10
        for scene in episode["scenes"]:
            for beat in scene["beats"]:
                ws.append([beat["fields"][header] for header in HEADERS])
                for cell in ws[ws.max_row]:
                    cell.font = Font(name="微软雅黑", size=10)
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    cell.border = border
        ws.freeze_panes = f"A{header_row + 1}"
        ws.auto_filter.ref = f"A{header_row}:M{ws.max_row}"
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def structural_findings(log: dict[str, Any]) -> list[dict[str, Any]]:
    findings = []
    unassigned = log.get("source_coverage", {}).get("unassigned_source_refs", [])
    if unassigned:
        findings.append({"code": "upstream_unassigned_source",
                         "message": "一级剧本仍有未映射到正文的来源段落，请先核对上游覆盖情况",
                         "beat_id": "", "source_refs": unassigned, "confidence": "high"})
    for episode in log["episodes"]:
        for scene in episode["scenes"]:
            for beat in scene["beats"]:
                if beat["source_origin"] == "generated":
                    findings.append({"code": "generated_source", "message": "该拍包含一级细纲创作内容，请核对创作依据", "beat_id": beat["id"], "source_refs": beat["source_refs"], "confidence": "high"})
                if any(not item.get("speaker") for item in beat["dialogue_blocks"]):
                    findings.append({"code": "missing_speaker", "message": "对白块未标明说话人，请核对一级剧本", "beat_id": beat["id"], "source_refs": beat["source_refs"], "confidence": "high"})
                for uncertainty in beat.get("uncertainties", []):
                    findings.append({"code": "creative_uncertainty", "message": str(uncertainty), "beat_id": beat["id"], "source_refs": beat["source_refs"], "confidence": "low"})
    return findings


def write_report(log: dict[str, Any], findings: list[dict[str, Any]], json_path: str | Path, md_path: str | Path) -> None:
    scene_count = sum(len(ep["scenes"]) for ep in log["episodes"])
    beat_count = sum(len(sc["beats"]) for ep in log["episodes"] for sc in ep["scenes"])
    report = {"schema_version": "scene_log_report_v1", "episodes": len(log["episodes"]),
              "scenes": scene_count, "beats": beat_count, "review_status": "pending_human_review",
              "findings": findings}
    Path(json_path).write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    lines = ["# 场记表审查报告", "", f"覆盖：{report['episodes']} 集、{scene_count} 场、{beat_count} 拍。", "",
             "摄影、灯光、表演与时长均为待审建议。疑点未自动修改剧本。", "", "## 疑点", ""]
    if findings:
        for finding in findings:
            lines.append(f"- **{finding.get('code', 'review')}** {finding.get('beat_id', '')}：{finding.get('message', '')}（来源：{', '.join(finding.get('source_refs', [])) or '待核对'}）")
    else:
        lines.append("- 本次未发现有文本依据的疑点。")
    Path(md_path).write_text("\n".join(lines) + "\n", encoding="utf-8")
