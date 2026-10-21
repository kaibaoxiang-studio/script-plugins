"""Cross-check scene-log JSON, workbook, report, and level-1 source."""

import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "ai-continuity-sheet" / "scripts"))
from scene_log_core import HEADERS, episode_script_text, load_source, prepare_scenes


def verify(source_path, output_dir):
    source = load_source(source_path)
    output = Path(output_dir)
    log = json.loads((output / "scene_log.json").read_text(encoding="utf-8"))
    report = json.loads((output / "scene_log_report.json").read_text(encoding="utf-8"))
    book = load_workbook(output / "scene_log.xlsx", read_only=False, data_only=True)
    by_id = {scene["scene_id"]: scene for scene in prepare_scenes(source)}
    by_number = {ep["number"]: ep for ep in source["episodes"]}
    assert len(book.sheetnames) == len(log["episodes"])
    total_scenes = total_beats = 0
    for episode in log["episodes"]:
        ws = book[f"EP{episode['number']:02d}"]
        source_script = episode_script_text(by_number[episode["number"]])
        header_row = next(row for row in range(1, ws.max_row + 1)
                          if [ws.cell(row, col).value for col in range(1, 14)] == HEADERS)
        chunks = []
        for row in range(1, header_row):
            value = ws.cell(row, 1).value
            assert isinstance(value, str)
            chunks.append(re.sub(r"^一级标准化剧本正文(?:（\d+）)?\n", "", value))
        assert "".join(chunks) == source_script
        excel_rows = [[ws.cell(row, col).value for col in range(1, 14)]
                      for row in range(header_row + 1, ws.max_row + 1)]
        json_rows = []
        for scene in episode["scenes"]:
            original = by_id[scene["id"]]
            flattened = [unit for beat in scene["beats"] for unit in beat["source_units"]]
            assert [unit["id"] for unit in flattened] == [f"u{index:03d}" for index in range(1, len(flattened) + 1)]
            assert scene["sc_number"] == original["sc_number"]
            blocks = by_number[episode["number"]]["scenes"]
            source_scene = next(item for item in blocks if item["id"] == scene["id"])
            for block_index, block in enumerate(source_scene["blocks"]):
                assert "".join(unit["text"] for unit in flattened if unit["block_index"] == block_index) == block["text"]
            source_dialogue = [block["text"] for block in source_scene["blocks"] if block["type"] == "dialogue"]
            result_dialogue = [item["text"] for beat in scene["beats"] for item in beat["dialogue_blocks"]]
            assert result_dialogue == source_dialogue
            for index, beat in enumerate(scene["beats"], 1):
                assert beat["segment_id"] == f"P{index:02d}"
                assert beat["beat_id"] == f"B{index:02d}"
                json_rows.append([beat["fields"][header] for header in HEADERS])
                total_beats += 1
            total_scenes += 1
        assert excel_rows == json_rows
    assert total_scenes == report["scenes"]
    assert total_beats == report["beats"]
    return {"status": "passed", "episodes": len(log["episodes"]),
            "scenes": total_scenes, "beats": total_beats, "findings": len(report["findings"])}


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source")
    parser.add_argument("output")
    args = parser.parse_args()
    print(json.dumps(verify(args.source, args.output), ensure_ascii=False))
