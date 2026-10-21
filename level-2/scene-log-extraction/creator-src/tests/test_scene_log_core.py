import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "ai-continuity-sheet" / "scripts"))
from scene_log_core import (HEADERS, assemble, load_source, prepare_scenes,
                            structural_findings, validate_groups, write_workbook, write_report)


def fixture(long=False):
    huge = "甲" * 33000 if long else "他打开盒子。她看见钥匙。"
    return {"schema_version": "standardized_screenplay_v1", "source_manifest": {},
            "episodes": [
                {"id": "ep-001", "number": 1, "title": "钥匙", "summary": "交接钥匙。",
                 "scenes": [{"id": "scn-001", "number": "1-1", "source_heading": "1-1 室内 内 日",
                             "location": "室内", "interior_exterior": "内", "time": "日",
                             "characters": ["甲", "乙"], "source_refs": ["L1"],
                             "blocks": [
                                 {"type": "action", "text": huge, "source_refs": ["L2"], "origin": "source"},
                                 {"type": "dialogue", "text": "你好。 Hello.", "speaker": "甲", "voice": "OS",
                                  "language": "mixed", "source_refs": ["L3"], "origin": "source"}]}]},
                {"id": "ep-002", "number": 2, "title": "归还", "summary": "归还钥匙。",
                 "scenes": [{"id": "scn-002", "number": "2-1", "source_heading": "2-1 门口 外 夜",
                             "location": "门口", "interior_exterior": "外", "time": "夜",
                             "characters": ["乙"], "source_refs": ["L4"],
                             "blocks": [
                                 {"type": "action", "text": "乙归还钥匙。", "source_refs": ["L5"], "origin": "generated"}]}]},
            ]}


def result_for(scene):
    groups = [{"unit_ids": [unit["id"]], "reason": "动作或台词推进"} for unit in scene["units"]]
    notes = []
    for group in groups:
        notes.append({"unit_ids": group["unit_ids"], "characters": list(scene["characters"]),
                      "props": ["钥匙"], "camera": "三脚架，50mm，眼平，静镜",
                      "lighting": "—", "duration_estimate": {"seconds": 5, "basis": "台词与动作估计"},
                      "micro_expression": "—", "action_performance": "—", "synopsis": "交接钥匙",
                      "rationale": "原文单元", "uncertainties": []})
    return {"segmentation": {"episode_id": scene["episode_id"],
                             "scene_id": scene["scene_id"], "groups": groups},
            "enrichment": {"beats": notes}}


class SceneLogCoreTest(unittest.TestCase):
    def test_source_fidelity_and_one_beat_per_p(self):
        source = fixture()
        scenes = prepare_scenes(source)
        log = assemble(source, scenes, {scene["scene_id"]: result_for(scene) for scene in scenes})
        self.assertEqual([scenes[0]["sc_number"], scenes[1]["sc_number"]], ["SC01", "SC02"])
        first = log["episodes"][0]["scenes"][0]["beats"]
        self.assertEqual([beat["segment_id"] for beat in first], ["P01", "P02", "P03"])
        self.assertEqual(first[-1]["dialogue_blocks"][0]["text"], "你好。 Hello.")
        self.assertEqual(first[-1]["fields"]["台词"], "甲（OS）：你好。 Hello.")
        self.assertEqual(log["episodes"][1]["scenes"][0]["beats"][0]["source_origin"], "generated")
        self.assertEqual(prepare_scenes(source, [2])[0]["sc_number"], "SC02")

    def test_reject_missing_or_duplicate_source_units(self):
        scene = prepare_scenes(fixture())[0]
        with self.assertRaisesRegex(ValueError, "完整覆盖"):
            validate_groups(scene, [{"unit_ids": ["u001", "u001"], "reason": "错误"}])

    def test_action_quote_kept_whole(self):
        source = fixture()
        source["episodes"][0]["scenes"][0]["blocks"][0]["text"] = "甲喊道：“回来！”随后转身。"
        units = prepare_scenes(source)[0]["units"]
        self.assertEqual([unit["text"] for unit in units[:2]], ["甲喊道：“回来！”", "随后转身。"])

    def test_missing_body_actionable(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.json"
            source = fixture()
            source["episodes"][0]["scenes"][0]["blocks"] = []
            path.write_text(json.dumps(source, ensure_ascii=False), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "请先补全一级剧本"):
                load_source(path)

    def test_reject_unbased_character_and_report_upstream_gap(self):
        source = fixture()
        source["coverage"] = {"unassigned_source_refs": ["L99"]}
        scenes = prepare_scenes(source)
        results = {scene["scene_id"]: result_for(scene) for scene in scenes}
        results[scenes[0]["scene_id"]]["enrichment"]["beats"][0]["characters"].append("陌生人")
        with self.assertRaisesRegex(ValueError, "没有场景或说话人依据"):
            assemble(source, scenes, results)
        results[scenes[0]["scene_id"]] = result_for(scenes[0])
        log = assemble(source, scenes, results)
        findings = structural_findings(log)
        self.assertEqual(findings[0]["code"], "upstream_unassigned_source")
        self.assertEqual(findings[0]["source_refs"], ["L99"])

    def test_long_text_and_two_episode_sheets(self):
        source = fixture(long=True)
        scenes = prepare_scenes(source)
        log = assemble(source, scenes, {scene["scene_id"]: result_for(scene) for scene in scenes})
        with tempfile.TemporaryDirectory() as directory:
            workbook = Path(directory) / "scene_log.xlsx"
            write_workbook(log, source, workbook)
            book = load_workbook(workbook)
            self.assertEqual(book.sheetnames, ["EP01", "EP02"])
            sheet = book["EP01"]
            self.assertGreater(len(sheet.merged_cells.ranges), 1)
            chunks = [sheet.cell(row, 1).value for row in range(1, 3)]
            self.assertIn("甲" * 100, "".join(chunks))
            self.assertIn(HEADERS, [[cell.value for cell in row] for row in sheet.iter_rows()])
            report_json = Path(directory) / "scene_log_report.json"
            report_md = Path(directory) / "scene_log_report.md"
            write_report(log, [], report_json, report_md)
            self.assertEqual(json.loads(report_json.read_text())["episodes"], 2)


if __name__ == "__main__":
    unittest.main()
